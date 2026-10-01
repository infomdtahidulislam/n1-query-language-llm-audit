#!/usr/bin/env python3
"""build_reapprove.py — N1: re-approval sheets for renderings whose text changed after round 1.

A rendering that was edited no longer carries its old approval, so its two approvers judge the new
text. Each sheet shows what changed (old word -> new word, or the whole previous line), the approver's
own earlier decision, and nothing about the other approver.

Edit CHANGED below, then run:  python3 build_reapprove.py kit.xlsx outdir
"""
import re, sys, os, json
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, Protection
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

KIT = sys.argv[1] if len(sys.argv) > 1 else "KIT-v2.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "."
ROUND = "R1"

# qid -> (what the approvers saw before, one line saying what happened)
CHANGED = {
 "Q185": ("dhakar bahire kon ecommercee site er delivery valo?",
          "You accepted this row in round 1. The authors asked the writer about one word; the writer said it was "
          "a typing slip and corrected it — 'ecommercee' is now 'ecommerce'. Nothing else changed. Judge the corrected line."),
 "Q223": ("Official nki unoffical laptop kena valo?",
          "You accepted this row in round 1. The authors asked the writer about one word; the writer said it was "
          "a typing slip and corrected it — 'unoffical' is now 'unofficial'. Nothing else changed. Judge the corrected line."),
}

prev = json.load(open("decisions.json")) if os.path.exists("decisions.json") else {}

wb = openpyxl.load_workbook(KIT, data_only=True)["Queries"]
hdr = [c.value for c in wb[1]]; ix = {h: i for i, h in enumerate(hdr)}
K = {}
for r in wb.iter_rows(min_row=2, values_only=True):
    q = str(r[ix["query_id"]] or "").strip()
    if re.fullmatch(r"Q\d{3}", q):
        bn = str(r[ix["bn_text"]] or "")
        K[q] = dict(cat=r[ix["category"]], bn=bn, bl=str(r[ix["bl_text"]] or ""),
                    au=str(r[ix["bl_author"]] or ""),
                    ap=[str(r[ix["bl_approve_1"]] or ""), str(r[ix["bl_approve_2"]] or "")],
                    digits=" ".join(re.findall(r"\d+(?:\.\d+)?", bn)))

hfill=PatternFill("solid",fgColor="1F3864"); yfill=PatternFill("solid",fgColor="FFF2CC")
bfill=PatternFill("solid",fgColor="EAF1FB"); gfill=PatternFill("solid",fgColor="E2EFDA")
thin=Side(style="thin",color="BFBFBF"); bord=Border(left=thin,right=thin,top=thin,bottom=thin)
BN="Nirmala UI"
HEAD=["query_id","category","digits","Bangla question","Banglish (corrected)","What changed",
      "Your earlier decision","Accept / Reject","Reason — required if you Reject"]
WIDTH=[10,15,12,42,42,52,18,15,42]
RULES=[("THE RULES HAVE NOT CHANGED",True),("",False),
 ("Spelling and word choice are NOT grounds for rejection. Meaning unchanged + a real person could type it = Accept.",False),
 ("বানান বা শব্দ পছন্দ না হলেই Reject নয়।",False),("",False),("REJECT only for these",True),("",False),
 ("1. A number changed, shortened or spelled out — the digits must match the Bangla exactly.",False),
 ("2. A brand or shop name appears.",False),("3. Transliteration-tool output.",False),
 ("4. Something in the Bangla is missing, or something added that the Bangla does not say.",False),
 ("5. Bangla letters in the Banglish, or more than two sentences.",False),
 ("6. No real person would type it — including a clear typo that makes a non-word.",False),("",False),
 ("Work alone. Do not discuss any row with the other raters. Reasons come to Tahidul.",False),
 ("Save the file and send it back. Do not rename it.",False)]

def build(code, qs):
    out=openpyxl.Workbook(); w=out.active; w.title="Judge"
    w["A1"]=f"N1 · Banglish re-approval — rater {code} — {len(qs)} row(s)"
    w["A1"].font=Font(bold=True,size=13,color="1F3864")
    w["A2"]=("These renderings were edited after you judged them, so the old decision no longer covers them. "
             "Judge the corrected line: Accept or Reject, with a reason on any Reject.")
    w["A2"].font=Font(size=10,italic=True,color="595959")
    w.merge_cells(start_row=1,end_row=1,start_column=1,end_column=len(HEAD))
    w.merge_cells(start_row=2,end_row=2,start_column=1,end_column=len(HEAD))
    for j,(h,wd) in enumerate(zip(HEAD,WIDTH),start=1):
        c=w.cell(row=4,column=j,value=h); c.font=Font(bold=True,color="FFFFFF",size=10); c.fill=hfill
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=bord
        w.column_dimensions[get_column_letter(j)].width=wd
    w.row_dimensions[4].height=30
    for i,q in enumerate(sorted(qs)):
        r=5+i; k=K[q]; old,note=CHANGED[q]
        p=prev.get(q,{}).get(code,("",""))
        ptxt=("Accept" if p[0]=="accept" else f"Reject — “{p[1]}”") if p and p[0] else "—"
        for j,v in enumerate([q,k["cat"],k["digits"],k["bn"],k["bl"],note,ptxt,None,None],start=1):
            c=w.cell(row=r,column=j,value=v); c.border=bord
            c.alignment=Alignment(vertical="top",wrap_text=(j in (4,5,6,7,9)))
            c.protection=Protection(locked=(j not in (8,9)))
            if j==1: c.font=Font(bold=True,size=10)
            if j==3: c.font=Font(bold=True,size=11,color="C00000"); c.alignment=Alignment(horizontal="center",vertical="top")
            if j==4: c.font=Font(name=BN,size=12)
            if j==5: c.fill=gfill; c.font=Font(size=11)
            if j==6: c.font=Font(size=10,italic=True,color="1F4E79")
            if j==7: c.font=Font(size=10,color="595959")
            if j in (8,9): c.fill=yfill
            if j==8: c.alignment=Alignment(horizontal="center",vertical="center")
    last=4+len(qs)
    dv=DataValidation(type="list",formula1='"Accept,Reject"',allow_blank=True,showDropDown=False)
    w.add_data_validation(dv); dv.add(f"H5:H{last}")
    w.conditional_formatting.add(f"I5:I{last}",
        FormulaRule(formula=['AND($H5="Reject",$I5="")'],fill=PatternFill("solid",fgColor="F8CBAD")))
    w.freeze_panes="A5"; w.protection.sheet=True; w.protection.enable()
    r2=out.create_sheet("Rules"); r2.column_dimensions["A"].width=118
    for i,(t,head) in enumerate(RULES,start=1):
        c=r2.cell(row=i,column=1,value=t); c.alignment=Alignment(vertical="top",wrap_text=True)
        c.font=Font(bold=True,size=12,color="1F3864") if head else Font(name=BN,size=11)
    r2.sheet_view.showGridLines=False
    p=os.path.join(OUT,f"RATER-{code}-REAPPROVE-{ROUND}.xlsx"); out.save(p); return p

by={}
for q in CHANGED:
    for a in K[q]["ap"]: by.setdefault(a,[]).append(q)
ok=True; cov={}
for code,qs in sorted(by.items()):
    p=build(code,qs)
    chk=openpyxl.load_workbook(p); ws=chk["Judge"]
    ids=[r[0] for r in ws.iter_rows(min_row=5,values_only=True) if r and r[0]]
    blob="\n".join(str(c.value) for s in chk.worksheets for row in s.iter_rows() for c in row if c.value)
    names=[t for t in ("R1","R2","R3") if t in blob]
    assigned=all(code in K[q]["ap"] for q in ids); notmine=all(K[q]["au"]!=code for q in ids)
    texts=all(K[q]["bl"]==r[4] for q,r in zip(ids,[r for r in ws.iter_rows(min_row=5,values_only=True) if r and r[0]]))
    for q in ids: cov[q]=cov.get(q,0)+1
    print(f"  {p}: {ids} | assigned to {code}: {assigned} | none self-written: {notmine} | "
          f"shows current kit text: {texts} | codes visible: {names}")
    ok &= assigned and notmine and texts and names==[code]
print("\ncoverage:",cov,"— each row twice:",all(v==2 for v in cov.values()))
print("VERDICT:","PASS" if ok and all(v==2 for v in cov.values()) else "FAIL")
