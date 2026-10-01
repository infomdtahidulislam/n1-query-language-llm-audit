#!/usr/bin/env python3
"""build_approver_files.py — N1: per-rater Banglish APPROVAL files, blind to en_text and to the writer.

Each rater judges the rows they did NOT write — the rows where the kit names them in
bl_approve_1 or bl_approve_2 (166 / 167 / 167; 500 judgements over 250 rows, two each).

Deliberately absent from every file:
  * en_text, in any form                     (A.4b blindness)
  * who wrote the rendering                  (so the judgement is on the text, not the person)
  * the other approver's decision            (A.5 independent approval)
Row order is shuffled with a fixed per-rater seed, so the two approvers of a row meet it at
different points and fatigue does not line up with category.

Usage: python3 build_approver_files.py kit.xlsx outdir
"""
import re, sys, os, random
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, Protection
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

KIT = sys.argv[1] if len(sys.argv) > 1 else "QUERY-AUTHORING-KIT.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "."
SEED = {"R1": 10101, "R2": 20202, "R3": 30303}          # recorded, so a file can be rebuilt identically

NOTE_FOR_RATER = {"Q135": "শরিয়াহভিত্তিক is deliberate. If the Banglish turns it into 'Islami Bank' / 'ইসলামী ব্যাংক', that is a brand name — reject."}
MONTHLY = "The amount in this question is a MONTHLY internet package fee, not a one-off purchase price."

wb = openpyxl.load_workbook(KIT, data_only=True)
ws = wb["Queries"]
hdr = [c.value for c in ws[1]]
ix = {h: i for i, h in enumerate(hdr)}
rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    q = str(r[ix["query_id"]] or "").strip()
    if not re.fullmatch(r"Q\d{3}", q):
        continue
    bn = str(r[ix["bn_text"]] or "").strip()
    note = str(r[ix["notes"]] or "").strip()
    if q in NOTE_FOR_RATER:        note = NOTE_FOR_RATER[q]
    elif "MONTHLY package fee" in note: note = MONTHLY
    rows.append(dict(id=q, cat=str(r[ix["category"]] or "").strip(), sub=str(r[ix["subtype"]] or "").strip(),
                     bn=bn, bl=str(r[ix["bl_text"]] or "").strip(), en=str(r[ix["en_text"]] or "").strip(),
                     digits=" ".join(re.findall(r"\d+(?:\.\d+)?", bn)), note=note,
                     au=str(r[ix["bl_author"]] or "").strip(),
                     ap=[str(r[ix["bl_approve_1"]] or "").strip(), str(r[ix["bl_approve_2"]] or "").strip()]))

HEAD  = ["query_id", "category", "type", "digits", "Bangla question", "Banglish to judge",
         "Accept / Reject", "Reason — required if you Reject", "note"]
WIDTH = [10, 14, 15, 12, 50, 50, 15, 46, 40]
BN_FONT = "Nirmala UI"
hfill = PatternFill("solid", fgColor="1F3864")
dfill = PatternFill("solid", fgColor="FFF2CC")     # the two columns they fill
jfill = PatternFill("solid", fgColor="EAF1FB")     # the text under judgement
nfill = PatternFill("solid", fgColor="F2F2F2")
thin  = Side(style="thin", color="BFBFBF")
bord  = Border(left=thin, right=thin, top=thin, bottom=thin)

RULES = [
 ("HOW TO USE THIS FILE", True), ("", False),
 ("For every row: read the Bangla, read the Banglish, then put Accept or Reject in the yellow column.", False),
 ("Reject → write a reason the writer can act on, in the next column. Accept needs no reason.", False),
 ("প্রতিটি row-এ বাংলা ও বাংলিশ পড়ে Accept বা Reject লিখুন। Reject করলে কারণ লিখুন।", False),
 ("", False),
 ("Work alone. Do not discuss any row with the other two raters. Questions go to Tahidul, never to them.", False),
 ("You are not told who wrote each row, on purpose. Judge the text, not the person.", False),
 ("একা কাজ করুন। কে লিখেছে তা জানানো হয়নি — লেখা দেখেই বিচার করুন।", False),
 ("", False),
 ("THE ONLY QUESTION", True), ("", False),
 ("“Could a real person have typed this?”  —  not  “is this how I would have written it?”", False),
 ("“এটা কি একজন সত্যিকারের মানুষ এভাবে টাইপ করতে পারে?”", False),
 ("", False),
 ("REJECT if any of these is true", True), ("", False),
 ("1. A number changed, or was shortened or spelled out. The digits must match the Bangla exactly —", False),
 ("    25000 stays 25000. Not 25,000, not 25k, not 'pochish hazar'. The 'digits' column shows what must appear.", False),
 ("2. A brand or shop name appears — Walton, Samsung, bKash, Daraz, and so on. Even as an example.", False),
 ("3. It looks like transliteration-tool output (Pānira philṭāra, br̥yāṇḍera). Nobody types that.", False),
 ("4. Something in the Bangla is missing, or something is there that the Bangla does not say —", False),
 ("    a dropped price limit, a dropped 'for a student', a 'best' where the Bangla only says 'good'.", False),
 ("5. Bangla letters appear in the Banglish, or it runs to more than two sentences.", False),
 ("6. No real person would type it — wrong word order, a word that does not exist, unreadable.", False),
 ("", False),
 ("DO NOT REJECT for any of these", True), ("", False),
 ("• Spelling. valo or bhalo, moddhe or modde, theke or thake — all correct. There is no standard.", False),
 ("• Capital or small letters, a hyphen, a comma, a missing question mark. None of these is a rule.", False),
 ("• A word you would not have chosen, when the meaning is the same and a real person could type it.", False),
 ("বানান, বড়/ছোট হাতের অক্ষর, হাইফেন, কমা — এগুলো কোনো নিয়ম নয়। এসবের জন্য Reject করবেন না।", False),
 ("", False),
 ("WHAT HAPPENS NEXT", True), ("", False),
 ("Both of the other raters judge every row independently. A rendering enters the study only if both", False),
 ("accept it. If either rejects, your reason goes to the writer, they revise, and you both judge it again —", False),
 ("without comparing notes. Three rounds at most, then the authors decide.", False),
 ("", False),
 ("When you are done: save the file and send it back. Do not rename it, do not add, delete or re-order rows.", False),
]

def build(code, mine, seed):
    random.Random(seed).shuffle(mine)
    out = openpyxl.Workbook()
    w = out.active; w.title = "Judge"
    w["A1"] = f"N1 · Banglish approval sheet — rater {code} — {len(mine)} rows to judge"
    w["A1"].font = Font(bold=True, size=13, color="1F3864")
    w["A2"] = ("Accept or Reject every row. Reject needs a reason. Read the 'Rules' tab first. "
               "Work alone — do not discuss any row with the other raters.")
    w["A2"].font = Font(size=10, italic=True, color="595959")
    w.merge_cells(start_row=1, end_row=1, start_column=1, end_column=9)
    w.merge_cells(start_row=2, end_row=2, start_column=1, end_column=9)
    for j, (h, wd) in enumerate(zip(HEAD, WIDTH), start=1):
        c = w.cell(row=4, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10); c.fill = hfill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = bord
        w.column_dimensions[get_column_letter(j)].width = wd
    w.row_dimensions[4].height = 30
    for i, x in enumerate(mine):
        r = 5 + i
        for j, v in enumerate([x["id"], x["cat"], x["sub"], x["digits"], x["bn"], x["bl"], None, None, x["note"]], start=1):
            c = w.cell(row=r, column=j, value=v)
            c.border = bord
            c.alignment = Alignment(vertical="top", wrap_text=(j in (5, 6, 8, 9)))
            c.protection = Protection(locked=(j not in (7, 8)))
            if j == 1: c.font = Font(bold=True, size=10)
            if j == 4:
                c.font = Font(bold=True, size=11, color="C00000")
                c.alignment = Alignment(horizontal="center", vertical="top")
            if j == 5: c.font = Font(name=BN_FONT, size=12)
            if j == 6: c.font = Font(size=11); c.fill = jfill
            if j in (7, 8): c.fill = dfill
            if j == 7: c.alignment = Alignment(horizontal="center", vertical="center")
            if j == 9: c.fill = nfill; c.font = Font(size=9, italic=True, color="7F6000")
    last = 4 + len(mine)
    dv = DataValidation(type="list", formula1='"Accept,Reject"', allow_blank=True, showDropDown=False,
                        errorTitle="Accept or Reject", error="Type Accept or Reject.")
    w.add_data_validation(dv); dv.add(f"G5:G{last}")
    # a Reject with no reason turns the reason cell red
    w.conditional_formatting.add(f"H5:H{last}",
        FormulaRule(formula=[f'AND($G5="Reject",$H5="")'], fill=PatternFill("solid", fgColor="F8CBAD")))
    w.freeze_panes = "A5"
    w.protection.sheet = True; w.protection.enable()

    r2 = out.create_sheet("Rules")
    r2.column_dimensions["A"].width = 122
    for i, (t, head) in enumerate(RULES, start=1):
        c = r2.cell(row=i, column=1, value=t)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.font = Font(bold=True, size=12, color="1F3864") if head else Font(name=BN_FONT, size=11)
    r2.sheet_view.showGridLines = False
    p = os.path.join(OUT, f"RATER-{code}-APPROVE.xlsx")
    out.save(p); return p

made = []
for code in ("R1", "R2", "R3"):
    mine = [dict(x) for x in rows if code in x["ap"]]
    made.append((code, build(code, mine, SEED[code]), len(mine)))

# ---------------- verification ----------------
kit = {x["id"]: x for x in rows}
all_en = [x["en"] for x in rows if x["en"]]
counts, ok = {}, True
for code, path, n in made:
    chk = openpyxl.load_workbook(path)
    jw = chk["Judge"]
    data = [r for r in jw.iter_rows(min_row=5, values_only=True) if r and r[0]]
    ids = [r[0] for r in data]
    self_written = [i for i in ids if kit[i]["au"] == code]
    bn_ok = all(kit[r[0]]["bn"] == r[4] for r in data)
    bl_ok = all(kit[r[0]]["bl"] == r[5] for r in data)
    empty = all(r[6] in (None, "") and r[7] in (None, "") for r in data)
    blob = "\n".join(str(c.value) for s in chk.worksheets for row in s.iter_rows() for c in row if c.value)
    leak = [e for e in all_en if e and e in blob]
    authorleak = re.findall(r"written by R[123]|author.{0,12}R[123]", blob)
    for i in ids: counts[i] = counts.get(i, 0) + 1
    print(f"{code}: {len(ids)} rows · none self-written {not self_written} · bn identical {bn_ok} · "
          f"bl identical {bl_ok} · decision cols empty {empty} · en_text found {len(leak)} · "
          f"writer identity leaks {len(authorleak)} · shuffled {ids[:3]}")
    ok &= not self_written and bn_ok and bl_ok and empty and not leak and not authorleak
twice = all(v == 2 for v in counts.values())
print(f"every row judged exactly twice: {twice} ({len(counts)} distinct rows, {sum(counts.values())} judgements)")
print("counts:", {c: n for c, _, n in made}, "| seeds:", SEED)
print("VERDICT:", "PASS" if ok and twice and len(counts) == 250 else "FAIL")
