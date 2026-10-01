#!/usr/bin/env python3
"""build_rater_export.py — N1: per-rater Banglish WRITING files, blind to en_text.

One workbook per rater, containing only the rows that rater writes (kit column bl_author).
Carries query_id / category / type / the digits that must be copied / bn_text / an empty
Banglish column / a rater-facing note.  en_text NEVER appears, in any form.

Usage:  python3 build_rater_export.py [kit.xlsx] [outdir]
"""
import re, sys, os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, Protection
from openpyxl.utils import get_column_letter

KIT = sys.argv[1] if len(sys.argv) > 1 else "QUERY-AUTHORING-KIT.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "."

# Kit notes are written for the authors and name kit columns. Re-phrase for a rater.
NOTE_FOR_RATER = {
    "Q135": "শরিয়াহভিত্তিক is deliberate. Do NOT turn it into 'Islami Bank' / 'ইসলামী ব্যাংক' — that is a bank name, and brand names are not allowed.",
}
MONTHLY = "The amount in this question is a MONTHLY internet package fee, not a one-off purchase price. Keep it that way."

wb = openpyxl.load_workbook(KIT, data_only=True)
ws = wb["Queries"]
hdr = [c.value for c in ws[1]]
ix = {h: i for i, h in enumerate(hdr)}
rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    qid = str(r[ix["query_id"]] or "").strip()
    if not re.fullmatch(r"Q\d{3}", qid):
        continue
    bn = str(r[ix["bn_text"]] or "").strip()
    note = str(r[ix["notes"]] or "").strip()
    if qid in NOTE_FOR_RATER:
        note = NOTE_FOR_RATER[qid]
    elif "MONTHLY package fee" in note:
        note = MONTHLY
    rows.append(dict(id=qid, cat=str(r[ix["category"]] or "").strip(),
                     sub=str(r[ix["subtype"]] or "").strip(), bn=bn,
                     en=str(r[ix["en_text"]] or "").strip(),
                     digits=" ".join(re.findall(r"\d+(?:\.\d+)?", bn)),
                     author=str(r[ix["bl_author"]] or "").strip(), note=note))

HEAD = ["query_id", "category", "type", "digits to copy exactly", "Bangla question",
        "BANGLISH — write here", "note"]
WIDTH = [10, 15, 16, 16, 58, 58, 46]
BN_FONT = "Nirmala UI"

hfill = PatternFill("solid", fgColor="1F3864")
wfill = PatternFill("solid", fgColor="FFF2CC")          # the column they type in
nfill = PatternFill("solid", fgColor="F2F2F2")
thin  = Side(style="thin", color="BFBFBF")
bord  = Border(left=thin, right=thin, top=thin, bottom=thin)

RULES = [
 ("HOW TO USE THIS FILE", True),
 ("", False),
 ("Type your Banglish into the yellow column only. Everything else is locked so it cannot be changed by accident.", False),
 ("যে হলুদ ঘরগুলো আছে সেখানেই শুধু আপনার বাংলিশ লিখুন। বাকি সব লক করা আছে।", False),
 ("Save the file and send it back. Do not rename it, do not add or delete rows, do not re-order them.", False),
 ("Work alone. Do not discuss any row with the other two raters while this phase is open.", False),
 ("একা কাজ করুন। এই ধাপ চলার সময় অন্য দুজনের সাথে কোনো row নিয়ে আলোচনা করবেন না।", False),
 ("", False),
 ("THE RULES  ·  নিয়ম", True),
 ("", False),
 ("1. Write it as you would really type it in a chat — Facebook, Messenger, WhatsApp. Not textbook, not formal.", False),
 ("    চ্যাটে আপনি যেভাবে সত্যিই লেখেন, সেভাবেই লিখুন।", False),
 ("2. English loanwords are welcome — budget, phone, service, warranty, delivery. People use them; keep them.", False),
 ("3. Never use a transliteration tool or converter. Type it yourself. Tool output does not read like a person.", False),
 ("    কোনো ট্রান্সলিটারেশন টুল ব্যবহার করবেন না — নিজে টাইপ করুন।", False),
 ("4. Numbers stay exactly as they are in the Bangla — 8000 stays 8000. No comma, no k, no words.", False),
 ("    Amounts are written without a thousands separator — 25000, not 25,000. Do not add one.", False),
 ("    টাকার অঙ্কে কমা ব্যবহার করা হয় না — 25000, 25,000 নয়। বাংলায় যা আছে হুবহু সেটাই লিখুন।", False),
 ("    The 'digits to copy exactly' column shows what must appear unchanged. If it is blank, the question has no number.", False),
 ("5. Never add a brand or shop name — no Walton, Samsung, bKash, Daraz, and so on. Not even as an example.", False),
 ("6. Keep the same meaning. Same question, same conditions. Do not drop a budget limit, do not add a detail that is not there.", False),
 ("7. Latin letters only — no Bangla script in your answer. One question, at most two sentences.", False),
 ("", False),
 ("SPELLING  ·  বানান", True),
 ("", False),
 ("There is no correct spelling in Banglish. valo or bhalo, moddhe or modde, theke or thake, capital or small — all fine.", False),
 ("Write it your way. The other raters spell things differently and that is expected; nobody will reject your row over spelling.", False),
 ("বাংলিশে কোনো নির্দিষ্ট বানান নেই। আপনার মতো করেই লিখুন।", False),
 ("", False),
 ("WHAT COMES NEXT", True),
 ("", False),
 ("When all three of you have finished writing, you will each receive a separate file with the other two raters' rows to", False),
 ("approve or reject — one row at a time, on your own, with a written reason for anything you reject. That is a later step.", False),
 ("The only question there is ever: “Could a real person have typed this?” — not “is this how I would have written it?”", False),
 ("“এটা কি একজন সত্যিকারের মানুষ এভাবে টাইপ করতে পারে?”", False),
 ("", False),
 ("Questions about a row go to Tahidul, not to the other raters.", False),
]

def build(code, myrows):
    out = openpyxl.Workbook()
    ws1 = out.active
    ws1.title = "Write"
    ws1["A1"] = f"N1 · Banglish writing sheet — rater {code} — {len(myrows)} rows"
    ws1["A1"].font = Font(bold=True, size=13, color="1F3864")
    ws1["A2"] = ("Type your Banglish in the yellow column. Read the 'Rules' tab first. "
                 "Work alone — do not discuss any row with the other raters.")
    ws1["A2"].font = Font(size=10, italic=True, color="595959")
    ws1.merge_cells(start_row=1, end_row=1, start_column=1, end_column=7)
    ws1.merge_cells(start_row=2, end_row=2, start_column=1, end_column=7)

    for j, (h, w) in enumerate(zip(HEAD, WIDTH), start=1):
        c = ws1.cell(row=4, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = hfill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = bord
        ws1.column_dimensions[get_column_letter(j)].width = w
    ws1.row_dimensions[4].height = 30

    for i, x in enumerate(myrows):
        r = 5 + i
        vals = [x["id"], x["cat"], x["sub"], x["digits"], x["bn"], None, x["note"]]
        for j, v in enumerate(vals, start=1):
            c = ws1.cell(row=r, column=j, value=v)
            c.border = bord
            c.alignment = Alignment(vertical="top", wrap_text=(j in (5, 6, 7)))
            c.protection = Protection(locked=(j != 6))
            if j == 1:
                c.font = Font(bold=True, size=10)
            if j == 4:
                c.font = Font(bold=True, size=11, color="C00000")
                c.alignment = Alignment(horizontal="center", vertical="top")
            if j == 5:
                c.font = Font(name=BN_FONT, size=12)
            if j == 6:
                c.fill = wfill
                c.font = Font(size=11)
            if j == 7:
                c.fill = nfill
                c.font = Font(size=9, italic=True, color="7F6000")
        # no explicit height: Excel auto-fits the wrapped rows, and the yellow cell grows as they type
    ws1.freeze_panes = "A5"
    ws1.protection.sheet = True          # guardrail, no password: only the yellow column is editable
    ws1.protection.enable()

    ws2 = out.create_sheet("Rules")
    ws2.column_dimensions["A"].width = 128
    for i, (t, is_head) in enumerate(RULES, start=1):
        c = ws2.cell(row=i, column=1, value=t)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if is_head:
            c.font = Font(bold=True, size=12, color="1F3864")
        else:
            c.font = Font(name=BN_FONT, size=11)
    ws2.sheet_view.showGridLines = False

    p = os.path.join(OUT, f"RATER-{code}-WRITE.xlsx")
    out.save(p)
    return p

made = []
for code in ("R1", "R2", "R3"):
    mine = sorted([x for x in rows if x["author"] == code], key=lambda z: z["id"])
    made.append((code, build(code, mine), len(mine)))

# ---------------- verification ----------------
print(f"kit rows: {len(rows)}")
all_en = [x["en"] for x in rows if x["en"]]
seen = set()
ok = True
for code, path, n in made:
    chk = openpyxl.load_workbook(path)
    w = chk["Write"]
    ids, texts = [], []
    for r in w.iter_rows(min_row=5, values_only=True):
        if not r[0]: continue
        ids.append(r[0]); texts.append(r)
    kit = {x["id"]: x for x in rows}
    bn_ok = all(kit[i]["bn"] == t[4] for i, t in zip(ids, texts))
    bl_empty = all(t[5] in (None, "") for t in texts)
    strings = [str(c.value) for s in chk.worksheets for row in s.iter_rows() for c in row if c.value]
    blob = "\n".join(strings)
    leak = [e for e in all_en if e and e in blob]
    author_ok = all(kit[i]["author"] == code for i in ids)
    dup = seen & set(ids); seen |= set(ids)
    print(f"{code}: {len(ids)} rows · bn identical {bn_ok} · banglish column empty {bl_empty} · "
          f"all rows assigned to {code} {author_ok} · en_text strings found {len(leak)} · overlap {len(dup)}")
    ok &= bn_ok and bl_empty and author_ok and not leak and not dup
print("union of all three:", len(seen), "expected 250 —", "OK" if len(seen) == 250 else "MISMATCH")
print("counts:", {c: n for c, _, n in made})
print("VERDICT:", "PASS" if ok and len(seen) == 250 else "FAIL")
