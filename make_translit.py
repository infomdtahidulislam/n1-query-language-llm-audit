#!/usr/bin/env python3
"""make_translit.py — generate the mechanical `bl_translit` rendering (Addendum §2 / prereg A.5b, F.10).

Reads QUERY-AUTHORING-KIT.xlsx, takes every row flagged robust50 = YES, romanizes bn_text with ONE fixed,
published, rule-based scheme, and writes the result. Nothing is hand-edited. Digits and punctuation pass through.

Usage:
  python make_translit.py QUERY-AUTHORING-KIT.xlsx                 # dry run: prints a table, writes bl_translit_preview.csv
  python make_translit.py QUERY-AUTHORING-KIT.xlsx --write         # also fills the bl_translit column in the workbook
  python make_translit.py QUERY-AUTHORING-KIT.xlsx --scheme HK     # try another scheme for the H.4 decision (default RomanColloquial)

Needs:  pip install openpyxl aksharamukha
Record in prereg H.4:  library name + version + scheme, printed at the end of every run.

Why Aksharamukha 'RomanColloquial' (tested 2 Sep 2026 on sample queries):
  - rule-based and deterministic, lowercase, no diacritics, no capitals inside words
  - maps ব to b (indic-transliteration's ITRANS/HK map it to v: "movAila", "vryAnDera") and handles য় cleanly
    (ITRANS/HK left a stray nukta: "DhAkAya়")
  - it is still NOT natural Banglish — it keeps the Sanskrit-style final vowel ("takara madhye kona mobaila bhalo habe")
    — and that is acceptable by design: bl_translit isolates SCRIPT with register held constant; naturalness is the job
    of the human bl arm. Say so in the paper in one sentence.
"""
import sys, csv
try:
    import openpyxl
    from aksharamukha import transliterate as ak
    import importlib.metadata as md
except ImportError as e:
    sys.exit(f"missing dependency: {e} — pip install openpyxl aksharamukha")

args = [a for a in sys.argv[1:] if not a.startswith("--")]
PATH = args[0] if args else "QUERY-AUTHORING-KIT.xlsx"
WRITE = "--write" in sys.argv
SCHEME = sys.argv[sys.argv.index("--scheme") + 1] if "--scheme" in sys.argv else "RomanColloquial"
VERSION = md.version("aksharamukha")

def romanize(bn: str) -> str:
    return ak.process("Bengali", SCHEME, bn)

wb = openpyxl.load_workbook(PATH)
ws = wb["Queries"]
hdr = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
col = {n: i for i, n in enumerate(hdr)}
for need in ("query_id", "bn_text", "robust50"):
    if need not in col: sys.exit(f"Queries sheet has no '{need}' column — add it first (see ADDENDUM §1–§2)")
if WRITE and "bl_translit" not in col: sys.exit("no 'bl_translit' column to write into — add it first")

rows, skipped = [], []
for r in ws.iter_rows(min_row=2):
    qid = r[col["query_id"]].value
    if not qid or str(r[col["robust50"]].value or "").strip().upper() != "YES": continue
    bn = str(r[col["bn_text"]].value or "").strip()
    if not bn: skipped.append(str(qid)); continue
    out = romanize(bn)
    rows.append((str(qid), bn, out))
    if WRITE: r[col["bl_translit"]].value = out

with open("bl_translit_preview.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["query_id", "bn_text", "bl_translit", "scheme", "library"])
    for qid, bn, out in rows: w.writerow([qid, bn, out, SCHEME, f"aksharamukha {VERSION}"])
if WRITE: wb.save(PATH)

for qid, bn, out in rows[:10]: print(f"{qid}  {bn}\n      -> {out}")
print(f"\n{len(rows)} robust50 rows romanized" + (f"; {len(skipped)} skipped (bn_text empty): {skipped[:8]}" if skipped else ""))
print(f"preview: bl_translit_preview.csv" + ("  |  workbook column bl_translit WRITTEN" if WRITE else "  |  dry run — add --write to fill the column"))
print(f"H.4 record:  aksharamukha {VERSION}, target scheme '{SCHEME}', source 'Bengali', no post-editing")
