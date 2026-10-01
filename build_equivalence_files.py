#!/usr/bin/env python3
"""build_equivalence_files.py — N1: the A.4b back-translation files, one per rater.

A.4b: each of the 50 robust50 rows is back-translated to English twice — once from the bn rendering by
the row's FIRST approver, once from the bl rendering by its SECOND approver — by someone who did not
write that row's Banglish, and blind to the original en_text. The assignment is already fixed in the
kit's Equivalence sheet; this reads it rather than inventing one, and exports it VALUES-ONLY, which is
what A.4b requires ("raters receive a values-only export of the Equivalence sheet, never the query file").

Deliberately absent from every file: en_text in any form, the other rendering of the same row, the other
rater's work, and the equivalence verdict — columns F/G/H of the Equivalence sheet are the AUTHORS' job,
because judging equivalence means comparing against en_text, which the raters must never see.

Usage: python3 build_equivalence_files.py kit.xlsx outdir
"""
import re, sys, os, random
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, Protection
from openpyxl.utils import get_column_letter

KIT = sys.argv[1] if len(sys.argv) > 1 else "KIT-v5.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "."
SEED = {"R1": 40401, "R2": 50502, "R3": 60603}
ARM  = {"bn": "Bangla", "bl": "Banglish"}

wv = openpyxl.load_workbook(KIT, data_only=True)      # cached values: the export must carry text, not formulas
eq, qs = wv["Equivalence"], wv["Queries"]
hdr = [c.value for c in qs[1]]; ix = {h: i for i, h in enumerate(hdr)}
K = {}
for r in qs.iter_rows(min_row=3, values_only=True):
    q = str(r[ix["query_id"]] or "").strip()
    if re.fullmatch(r"Q\d{3}", q):
        K[q] = dict(en=str(r[ix["en_text"]] or ""), au=str(r[ix["bl_author"]] or ""),
                    cat=str(r[ix["category"]] or ""))
tasks = []
for r in eq.iter_rows(min_row=2, max_row=101, values_only=True):
    if not r[0]: continue
    tasks.append(dict(id=str(r[0]).strip(), arm=str(r[1]).strip(), src=str(r[2] or "").strip(),
                      who=str(r[3] or "").strip()))
assert len(tasks) == 100, len(tasks)
assert all(t["src"] for t in tasks), "a source_text cell came through empty — recalc the workbook first"

HEAD  = ["task", "query_id", "written in", "The sentence", "ENGLISH TRANSLATION — write here",
         "Note (only if something is genuinely unclear)"]
WIDTH = [7, 10, 12, 52, 56, 40]
BN="Nirmala UI"
hfill=PatternFill("solid",fgColor="1F3864"); yfill=PatternFill("solid",fgColor="FFF2CC")
bnf=PatternFill("solid",fgColor="EAF1FB"); blf=PatternFill("solid",fgColor="F2F7EC")
thin=Side(style="thin",color="BFBFBF"); bord=Border(left=thin,right=thin,top=thin,bottom=thin)

HOWTO=[("WHAT WE ARE ASKING",True),("",False),
 ("Each row has one Bangla or Banglish sentence. Write what it says in English. That is the whole task.",False),
 ("প্রতিটি row-এ একটি বাংলা বা বাংলিশ বাক্য আছে। সেটি ইংরেজিতে যা বলছে, তাই লিখুন।",False),
 ("",False),
 ("WHY — read this, it changes how you should translate",True),("",False),
 ("We are checking whether the Bangla and the Banglish versions of a question really carry the same",False),
 ("meaning. Your translation is the measuring instrument. So translate what is ACTUALLY WRITTEN —",False),
 ("not what you think the question was probably meant to be.",False),
 ("If the sentence is missing something, your translation should be missing it too. If it is vague, your",False),
 ("translation should be vague. Do not repair it, do not complete it, do not improve the English.",False),
 ("যা লেখা আছে ঠিক তাই অনুবাদ করুন — যা লেখা উচিত ছিল তা নয়। ভুল বা অস্পষ্টতা ঠিক করবেন না।",False),
 ("",False),
 ("RULES",True),("",False),
 ("1. Keep every number exactly as it appears — 25000 stays 25000. Never 25,000, never 25k, never in words.",False),
 ("2. Plain everyday English. One sentence. It does not have to be elegant.",False),
 ("3. Do not add a brand or shop name, even if you think you know which one is meant.",False),
 ("4. Translate only the sentence in front of you. Do not try to recall or reconstruct an English original —",False),
 ("    you have never seen one, and guessing at it would destroy the whole check.",False),
 ("5. If something is genuinely unclear, translate it as best you can AND say so in the Note column.",False),
 ("    A flagged uncertainty is useful to us. A confident guess is not.",False),
 ("",False),
 ("Work alone. Do not discuss any row with the other raters. Questions come to Tahidul.",False),
 ("Nobody is given both versions of the same question, and nobody translates a row they wrote themselves.",False),
 ("Save the file and send it back. Do not rename it, do not add, delete or re-order rows.",False)]

made=[]
for code in sorted({t["who"] for t in tasks}):
    mine=[dict(t) for t in tasks if t["who"]==code]
    random.Random(SEED[code]).shuffle(mine)
    out=openpyxl.Workbook(); w=out.active; w.title="Translate"
    w["A1"]=f"N1 · Back-translation — rater {code} — {len(mine)} sentences"
    w["A1"].font=Font(bold=True,size=13,color="1F3864")
    w["A2"]=("Write what each sentence says in English. Translate what is written, not what it should have said. "
             "Read the 'How to' tab first — it explains why that matters.")
    w["A2"].font=Font(size=10,italic=True,color="595959")
    w.merge_cells(start_row=1,end_row=1,start_column=1,end_column=len(HEAD))
    w.merge_cells(start_row=2,end_row=2,start_column=1,end_column=len(HEAD))
    for j,(hh,wd) in enumerate(zip(HEAD,WIDTH),start=1):
        c=w.cell(row=4,column=j,value=hh); c.font=Font(bold=True,color="FFFFFF",size=10); c.fill=hfill
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=bord
        w.column_dimensions[get_column_letter(j)].width=wd
    w.row_dimensions[4].height=32
    for i,t in enumerate(mine):
        r=5+i
        for j,v in enumerate([i+1,t["id"],ARM[t["arm"]],t["src"],None,None],start=1):
            c=w.cell(row=r,column=j,value=v); c.border=bord
            c.alignment=Alignment(vertical="top",wrap_text=(j in (4,5,6)))
            c.protection=Protection(locked=(j not in (5,6)))
            if j==2: c.font=Font(bold=True,size=10)
            if j==3: c.font=Font(size=9,bold=True,color=("1F4E79" if t["arm"]=="bn" else "375623"))
            if j==4:
                c.fill=(bnf if t["arm"]=="bn" else blf)
                c.font=(Font(name=BN,size=12) if t["arm"]=="bn" else Font(size=11))
            if j in (5,6): c.fill=yfill
    w.freeze_panes="A5"; w.protection.sheet=True; w.protection.enable()
    r2=out.create_sheet("How to"); r2.column_dimensions["A"].width=112
    for i,(t_,head) in enumerate(HOWTO,start=1):
        c=r2.cell(row=i,column=1,value=t_); c.alignment=Alignment(vertical="top",wrap_text=True)
        c.font=Font(bold=True,size=12,color="1F3864") if head else Font(name=BN,size=11)
    r2.sheet_view.showGridLines=False
    p=os.path.join(OUT,f"EQUIV-{code}.xlsx"); out.save(p); made.append((p,code,len(mine)))

# ---------------- verification ----------------
print("built:")
seen=set(); ok=True
src_of={(t["id"],t["arm"]):t["src"] for t in tasks}
all_en=[K[q]["en"] for q in K if K[q]["en"]]
for p,code,n in made:
    chk=openpyxl.load_workbook(p); ws=chk["Translate"]
    data=[r for r in ws.iter_rows(min_row=5,values_only=True) if r and r[1]]
    pairs=[(r[1],"bn" if r[2]=="Bangla" else "bl") for r in data]
    blob="\n".join(str(c.value) for s in chk.worksheets for row in s.iter_rows() for c in row if c.value)
    leak=[e for e in all_en if e and e in blob]
    names=[x for x in ("R1","R2","R3") if x in blob]
    right=all(t["who"]==code for t in tasks if (t["id"],t["arm"]) in pairs)
    notmine=all(K[q]["au"]!=code for q,_ in pairs)
    srcok=all(r[3]==src_of[(r[1],"bn" if r[2]=="Bangla" else "bl")] for r in data)
    empty=all(r[4] in (None,"") and r[5] in (None,"") for r in data)
    onearm=len({q for q,_ in pairs})==len(pairs)
    seen|=set(pairs)
    print(f"  {p}: {n} tasks ({sum(1 for _,a in pairs if a=='bn')} bn / {sum(1 for _,a in pairs if a=='bl')} bl) | "
          f"all assigned to {code}: {right} | none self-written: {notmine} | source matches the kit: {srcok} | "
          f"answer cols empty: {empty} | never both arms of one row: {onearm} | en_text found: {len(leak)} | codes visible: {names}")
    ok &= right and notmine and srcok and empty and onearm and not leak and names==[code]
print(f"\ncoverage: {len(seen)}/100 tasks, {len({q for q,_ in seen})} rows × 2 arms — complete: {len(seen)==100}")
print("VERDICT:", "PASS" if ok and len(seen)==100 else "FAIL")
