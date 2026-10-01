#!/usr/bin/env python3
"""n1_pipeline.py — N1 study runner: call planner, cache-first caller, deterministic outcome coder,
extractor driver, cost ledger and run-window log.

    ####################################################################################
    #  SAFETY: THIS SCRIPT MAKES NO NETWORK CALL UNLESS --go IS PASSED **AND** THE      #
    #  PREFLIGHT PASSES. Without --go every command is a dry run that only reads files. #
    ####################################################################################

Design contract (from PREREGISTRATION-DRAFT.md v0.9 and PROJECT-BRIEF.md):
  * models.json is the SINGLE SOURCE OF TRUTH for model ids, temperature, max_tokens and the
    no-system-prompt rule. This file hardcodes no model id and no sampling parameter.   (C.2, C.3)
  * The API key is read from the GATEWAY_KEY environment variable only. It is never written
    to disk, never logged, never included in a cached record.                            (C.2)
  * Every call is cached and keyed on (phase, model id, query_id, arm, rep, draw, prompt hash,
    parameters). A key that already holds a SUCCESSFUL response is never re-called, so runs are
    resumable and re-runs cost nothing. Failures are stored separately so they do retry. (H.1)
  * Each record stores the raw response JSON plus finish_reason, the full usage object, the
    provider's `model` string, any fingerprint field, HTTP status, attempt history and UTC
    timestamps.                                                                     (C.2, C.7)
  * Outcome coding is a SEPARATE deterministic pass over stored responses (no model involved).
    The redraw rule (degenerate only, at most 2 redraws at shifted draw id, content never
    repaired) is applied there and every redraw is logged.                          (D.5, D.6)
  * Extraction is a separate cached step: primary extractor at temperature 0 with JSON-schema
    validation, cross-extractor over a stratified subsample.                        (E.1, E.2)
  * A per-model cost ledger, a --dry-run call plan with a cost estimate, a --smoke mode, and a
    run-window log (start, end, observed model strings) for C.7/H.4.

Commands
  export    kit (.xlsx) -> queries.csv of FINAL rows + queries.sha256                   [offline]
  plan      print the call plan and cost estimate; alias for `run --dry-run`            [offline]
  run       execute subject calls (needs --go; --smoke/--pilot/--main/--translit)
  code      deterministic outcome coding over cached responses; emits redraw plan       [offline]
  redraw    execute the redraw plan produced by `code` (needs --go)
  extract   run the extractors over coded answers (needs --go; --primary/--cross)
  finalize  fold the primary extraction into final outcome codes + adjudication list    [offline]
  aliases   export brand_aliases from the kit (--starter for smoke, --frozen for F2)     [offline]
  release   copy runs/ with the gateway URL redacted for review/release (H.5)           [offline]
  ledger    per-model token/cost ledger from the cache                                  [offline]
  window    run-window report and model-string drift check for C.7/H.4                  [offline]
  selftest  offline end-to-end test on synthetic responses; no network, no real data     [offline]

Every [PROPOSED] marker in this file is a threshold or judgement that must be approved by the
authors and written into the pre-registration BEFORE the freeze. See RUNNER-NOTES.md.
"""
from __future__ import annotations

import argparse, csv, hashlib, json, os, random, re, sys, time, unicodedata
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from threading import Lock
from urllib import request as urlrequest, error as urlerror

SCHEMA = "n1-call-1"
ARMS = ("en", "bn", "bl")                 # the three registered arms (B.2)
ARM_TRANSLIT = "bl_translit"              # labeled robustness arm (A.5b, F.10)
ARM_COLUMN = {"en": "en_text", "bn": "bn_text", "bl": "bl_text", ARM_TRANSLIT: "bl_translit"}
MAX_DRAWS = 3                             # 1 original + at most 2 redraws (D.6)

# ---- [PROPOSED] script thresholds (prereg D.4) — approved on smoke-test data, see `code` summary --
BANGLA_RATIO_BN = 0.60                    # D.4 [PROPOSED]
BANGLA_RATIO_LATIN = 0.20                 # D.4 [PROPOSED]
# ---- degeneracy rules — APPROVED by the authors 2 Sep 2026, written into prereg D.5 ------------
DEGEN_MIN_CHARS = 1                       # empty/whitespace-only content = degenerate
DEGEN_LOOP_REPEATS = 5                    # loop detector: minimum repeats of a unit
DEGEN_LOOP_MAX_PERIOD = 60                # ... longest repeating unit considered (chars)
DEGEN_LOOP_MIN_SPAN = 20                  # ... and the loop must span >= this many chars
                                          #   (so a single repeated character needs 20 of them,
                                          #    which keeps "!!!!!" or "...." out of the detector)
DEGEN_REPLACEMENT_RATIO = 0.10            # >10% U+FFFD / control chars = degenerate
                                          # truncation (finish_reason=length) is never degenerate by itself
# ---- [PROPOSED] dry-run cost assumptions, used only when the cache has no measurement --------
FALLBACK_PROMPT_TOKENS = 40               # [PROPOSED] per subject call
FALLBACK_COMPLETION_TOKENS = 700          # [PROPOSED] per subject call (reasoning models burn more)
MEASURE_MIN_CALLS = 5                     # use measured means once a model has this many calls

LIVE_ALLOWED = False                      # flipped to True only after preflight passes
_PRINT_LOCK = Lock()
_WINDOW_LOCK = Lock()


# ────────────────────────────────────────────────────────────────────── helpers
def now_iso() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.%f")[:-3] + "Z"


def sha256_text(s: str) -> str:
    return hashlib.sha256(s.encode("utf-8")).hexdigest()


def sha256_file(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def canonical(obj) -> str:
    return json.dumps(obj, sort_keys=True, ensure_ascii=False, separators=(",", ":"))


def say(*a):
    with _PRINT_LOCK:
        print(*a, flush=True)


def die(msg: str, code: int = 2):
    print(f"\nREFUSED: {msg}\n", file=sys.stderr)
    sys.exit(code)


def slug(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", s)


def append_jsonl(path: Path, rec: dict):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "a", encoding="utf-8") as f:
        f.write(json.dumps(rec, ensure_ascii=False) + "\n")


def read_jsonl(path: Path):
    if not path.exists():
        return []
    out = []
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                out.append(json.loads(line))
    return out


# ────────────────────────────────────────────────────────────── config / inputs
class Config:
    """models.json is authoritative: ids, temperatures, max_tokens, the no-system-prompt rule."""

    def __init__(self, path: Path):
        if not path.exists():
            die(f"{path} not found — models.json is the single source of truth and must be present")
        self.raw = json.loads(path.read_text(encoding="utf-8"))
        self.path = path
        self.sha256 = sha256_file(path)
        self.base = (self.raw.get("gateway_base") or "").rstrip("/")
        if not self.base:
            die("models.json has no gateway_base")
        # Per-ROLE routing (added 11 Sep 2026). Subjects are the object of study and stay on the
        # gateway their answers were collected through; the extractor is the measuring INSTRUMENT and
        # may sit on a different route, provided EVERY extraction uses the same one. Both routes are
        # recorded in models.json, hashed at freeze (H.4) and redacted for review (H.5).
        self.extractor_base = (self.raw.get("extractor_gateway_base") or self.base).rstrip("/")
        self.key_env = self.raw.get("key_env") or "GATEWAY_KEY"
        self.extractor_key_env = self.raw.get("extractor_key_env") or self.key_env
        d = self.raw.get("defaults", {})
        self.subject_defaults = d.get("subjects", {})
        self.extractor_defaults = d.get("extractors", {})
        if "temperature" not in self.subject_defaults or "max_tokens" not in self.subject_defaults:
            die("models.json defaults.subjects must define temperature and max_tokens")
        if self.subject_defaults.get("system_prompt", "MISSING") is not None:
            die("models.json defaults.subjects.system_prompt must be null (bare-prompt protocol, B.3)")
        self.subjects = self.raw.get("subjects", [])
        if not self.subjects:
            die("models.json lists no subjects")
        self.extractors = self.raw.get("extractors", {})

    def subject_params(self, subj: dict) -> dict:
        """Sampling parameters for one subject: defaults, overridden only by an explicit
        per-model key in models.json (e.g. a model that accepts only one temperature)."""
        p = {"temperature": float(self.subject_defaults["temperature"]),
             "max_tokens": int(self.subject_defaults["max_tokens"])}
        for k in ("temperature", "max_tokens"):
            if k in subj:
                p[k] = type(p[k])(subj[k])
        return p

    def route(self, role: str) -> tuple[str, str]:
        """(base_url, key_env_var) for this role. 'extractor' may differ from everything else."""
        if role == "extractor":
            return self.extractor_base, self.extractor_key_env
        return self.base, self.key_env

    # Keys in defaults.extractors that are NOT request parameters. `response_format` configures the
    # call and is added by build_body; the rest are human documentation. A doc string reaching this
    # dict would be sent to the API as a parameter AND folded into the cache key, so editing a
    # comment would invalidate every cached extraction — exclude them by name and by leading "_".
    _NON_PARAM_EXTRACTOR_KEYS = {"response_format", "note", "notes", "comment", "description",
                                 "rationale", "doc", "docs"}

    def extractor_params(self, role: str = "primary") -> dict:
        """temperature and max_tokens, plus any other key in defaults.extractors — so `reasoning`
        and `provider` come from models.json instead of being unreachable. They are part of the
        cache key by design: a different reasoning or provider setting is a different instrument
        and must not silently reuse an older record.

        Per-role overrides (G.2): extractors.<role>.params in models.json is merged over the
        defaults — a contest extractor is a different instrument (its own provider pin, its own
        reasoning setting) and gets its own cache keys through exactly this mechanism. An override
        value of null REMOVES the key (e.g. "provider": null lifts the pin for a probe)."""
        p = {"temperature": float(self.extractor_defaults.get("temperature", 0)),
             "max_tokens": int(self.extractor_defaults.get("max_tokens",
                                                           self.subject_defaults["max_tokens"]))}
        for k, v in self.extractor_defaults.items():
            if k in ("temperature", "max_tokens") or k in self._NON_PARAM_EXTRACTOR_KEYS:
                continue
            if k in ("primary", "cross_check", "fallback"):      # role entries, not parameters
                continue
            if k.startswith("_"):
                continue
            p[k] = v
        over = (self.extractors.get(role) or {}).get("params") or {}
        for k, v in over.items():
            if k == "response_format":
                # per-role response_format IS a param here (unlike the default, which build_body
                # adds outside the key): a role that changes or omits provider-side JSON
                # enforcement is a different instrument and must not share cache keys. None means
                # "deliberately omitted" and is kept in the dict so build_body can see it.
                p[k] = v
                continue
            if k in self._NON_PARAM_EXTRACTOR_KEYS or k.startswith("_"):
                continue
            if k == "temperature":
                p[k] = float(v)
            elif k == "max_tokens":
                p[k] = int(v)
            else:
                p[k] = v
        return {k: v for k, v in p.items() if v is not None or k == "response_format"}

    def extractor(self, role: str) -> dict:
        e = self.extractors.get(role)
        if not e:
            die(f"models.json has no extractors.{role}")
        return e


def load_queries(path: Path) -> list[dict]:
    if not path.exists():
        die(f"{path} not found — run `export` first (needs FINAL rows in the kit)")
    rows = list(csv.DictReader(open(path, encoding="utf-8-sig")))
    if not rows:
        die(f"{path} has no rows")
    return rows


def cmd_export(a):
    """Kit -> queries.csv (FINAL rows only) + queries.sha256. Offline."""
    try:
        import openpyxl
    except ImportError:
        die("pip install openpyxl")
    kit = Path(a.kit)
    if not kit.exists():
        die(f"{kit} not found")
    wb = openpyxl.load_workbook(kit, data_only=True)
    ws = wb["Queries"]
    hdr = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    idx = {n: i for i, n in enumerate(hdr)}
    for need in ("query_id", "category", "subtype", "en_text", "bn_text", "bl_text", "status"):
        if need not in idx:
            die(f"kit Queries sheet has no '{need}' column")
    cols = ["query_id", "category", "subtype", "budget_bdt", "en_text", "bn_text", "bl_text",
            "bl_translit", "calibration", "smoke_test", "robust50"]
    out, skipped = [], 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        qid = r[idx["query_id"]]
        if not qid or not re.fullmatch(r"Q\d{3}", str(qid).strip()):
            continue                                     # skips the EXAMPLE row
        if str(r[idx["status"]] or "").strip().lower() != "final":
            skipped += 1
            continue
        rec = {}
        for c in cols:
            v = r[idx[c]] if c in idx and idx[c] < len(r) else None
            rec[c] = "" if v is None else str(v).strip()
        if rec["budget_bdt"] in ("—", "-"):
            rec["budget_bdt"] = ""
        out.append(rec)
    if not out:
        die(f"no rows with status=final in {kit} — nothing to export "
            f"({skipped} rows are not final yet)")
    dest = Path(a.queries)
    with open(dest, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(out)
    digest = sha256_file(dest)
    Path(str(dest) + ".sha256").write_text(f"{digest}  {dest.name}\n", encoding="utf-8")
    say(f"exported {len(out)} FINAL rows -> {dest}   ({skipped} not final, skipped)")
    say(f"sha256 {digest}   (record in prereg H.4 / F1)")
    miss = [r["query_id"] for r in out if r["robust50"].upper() == "YES" and not r["bl_translit"]]
    if miss:
        say(f"note: {len(miss)} robust50 rows have no bl_translit yet — run make_translit.py: {miss[:8]}")


# ─────────────────────────────────────────────────────────────────────── cache
class Cache:
    """Cache-first store. A key that holds a SUCCESSFUL response is never re-called (H.1).
    Failures live under errors/ with the same key, so they retry on the next run."""

    def __init__(self, root: Path):
        self.root = root
        self.ok_dir = root / "responses"
        self.err_dir = root / "errors"

    @staticmethod
    def key(spec: dict, params: dict) -> str:
        # prompt hash is part of the key: editing a query text can never silently reuse an answer
        return sha256_text(canonical({
            "phase": spec["phase"], "model_id": spec["model_id"], "query_id": spec["query_id"],
            "arm": spec["arm"], "rep": spec["rep"], "draw": spec["draw"],
            "prompt_sha256": spec["prompt_sha256"], "params": params,
        }))[:32]

    def path(self, spec: dict, k: str, err: bool = False) -> Path:
        base = self.err_dir if err else self.ok_dir
        return (base / spec["phase"] / slug(spec["model_id"]) /
                f"{spec['query_id']}_{spec['arm']}_r{spec['rep']}d{spec['draw']}_{k[:12]}.json")

    def has(self, spec: dict, k: str) -> bool:
        return self.path(spec, k).exists()

    def write(self, spec: dict, k: str, rec: dict, err: bool = False):
        p = self.path(spec, k, err)
        p.parent.mkdir(parents=True, exist_ok=True)
        tmp = p.with_suffix(".tmp")
        tmp.write_text(json.dumps(rec, ensure_ascii=False, indent=1), encoding="utf-8")
        tmp.replace(p)                       # atomic: a killed run never leaves a half record
        return p

    def records(self, phase: str):
        d = self.ok_dir / phase
        if not d.exists():
            return
        for p in sorted(d.rglob("*.json")):
            try:
                yield json.loads(p.read_text(encoding="utf-8"))
            except json.JSONDecodeError:
                say(f"WARNING: unreadable cache record {p}")


# ──────────────────────────────────────────────────────────────────── planning
def build_subject_plan(cfg: Config, queries: list[dict], arms: tuple, reps: int,
                       phase: str, seed: int) -> list[dict]:
    """One spec per (query, arm, subject, rep). Randomised interleaved order per B.5."""
    specs = []
    for q in queries:
        for arm in arms:
            text = (q.get(ARM_COLUMN[arm]) or "").strip()
            if not text:
                continue                      # not written yet / no translit for this row
            for subj in cfg.subjects:
                for rep in range(1, reps + 1):
                    specs.append({
                        "phase": phase, "role": "subject",
                        "model_id": subj["id"], "model_name": subj.get("name", subj["id"]),
                        "query_id": q["query_id"], "category": q.get("category", ""),
                        "arm": arm, "rep": rep, "draw": 1,
                        "prompt": text, "prompt_sha256": sha256_text(text),
                    })
    rng = random.Random(seed)
    by_query = defaultdict(list)
    for s in specs:
        by_query[s["query_id"]].append(s)
    order = sorted(by_query)
    rng.shuffle(order)                        # randomised query order
    out = []
    for qid in order:
        cell = by_query[qid]
        rng.shuffle(cell)                     # interleave arms × models × reps within the query
        out.extend(cell)
    return out


def measured_means(cache: Cache, phase_hint: list[str]) -> dict:
    """Mean prompt/completion tokens per model from whatever the cache already holds."""
    agg = defaultdict(lambda: {"n": 0, "pt": 0, "ct": 0})
    for phase in phase_hint:
        for rec in cache.records(phase):
            u = (rec.get("meta") or {}).get("usage") or {}
            m = rec["spec"]["model_id"]
            if u.get("prompt_tokens") is None:
                continue
            a = agg[m]
            a["n"] += 1
            a["pt"] += int(u.get("prompt_tokens") or 0)
            a["ct"] += int(u.get("completion_tokens") or 0)
    return {m: {"n": a["n"], "pt": a["pt"] / a["n"], "ct": a["ct"] / a["n"]}
            for m, a in agg.items() if a["n"] >= MEASURE_MIN_CALLS}


def load_prices(path: Path) -> dict:
    """Optional prices.json: {"<model id>": {"in_per_1m": 0.15, "out_per_1m": 0.60}}"""
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        say(f"WARNING: {path} is not valid JSON — ignoring prices")
        return {}


def print_plan(cfg: Config, plan: list[dict], cache: Cache, prices: dict, means: dict,
               title: str) -> dict:
    todo = [s for s in plan if not cache.has(s, Cache.key(s, cfg.subject_params(
        next(x for x in cfg.subjects if x["id"] == s["model_id"]))))]
    cached = len(plan) - len(todo)
    say("")
    say("=" * 78)
    say(f"CALL PLAN — {title}")
    say("=" * 78)
    per_model, per_arm = Counter(), Counter()
    for s in todo:
        per_model[s["model_id"]] += 1
        per_arm[s["arm"]] += 1
    say(f"queries in plan   : {len({s['query_id'] for s in plan})}")
    say(f"arms              : {', '.join(sorted({s['arm'] for s in plan}))}"
        f"   ({', '.join(f'{a}={n}' for a, n in sorted(per_arm.items()))} calls to make)")
    say(f"calls in plan     : {len(plan)}")
    say(f"already cached    : {cached}   (free, will be skipped)")
    say(f"calls TO MAKE     : {len(todo)}")
    say("")
    say(f"{'model':<22}{'calls':>7}{'in tok':>11}{'out tok':>11}{'est USD':>10}  basis")
    total = 0.0
    est_rows = {}
    for subj in cfg.subjects:
        mid = subj["id"]
        n = per_model.get(mid, 0)
        if not n:
            continue
        m = means.get(mid)
        pt, ct = (m["pt"], m["ct"]) if m else (FALLBACK_PROMPT_TOKENS, FALLBACK_COMPLETION_TOKENS)
        basis = f"measured n={m['n']}" if m else "ASSUMPTION [PROPOSED]"
        pr = prices.get(mid) or {}
        usd = (pt * n / 1e6) * float(pr.get("in_per_1m", 0)) + \
              (ct * n / 1e6) * float(pr.get("out_per_1m", 0))
        if not pr:
            basis += " · no price"
        total += usd
        est_rows[mid] = {"calls": n, "in_tok": pt * n, "out_tok": ct * n, "usd": usd,
                         "basis": basis}
        say(f"{subj.get('name', mid):<22}{n:>7}{pt*n:>11,.0f}{ct*n:>11,.0f}{usd:>10,.2f}  {basis}")
    say("-" * 78)
    say(f"{'TOTAL':<22}{len(todo):>7}"
        f"{sum(r['in_tok'] for r in est_rows.values()):>11,.0f}"
        f"{sum(r['out_tok'] for r in est_rows.values()):>11,.0f}{total:>10,.2f}")
    if not prices:
        say("no prices.json found -> USD column is 0.00; token counts are the real estimate.")
    if any("ASSUMPTION" in r["basis"] for r in est_rows.values()):
        say(f"ASSUMPTION rows use [PROPOSED] {FALLBACK_PROMPT_TOKENS} in / "
            f"{FALLBACK_COMPLETION_TOKENS} out tokens per call; the smoke test replaces them "
            f"with metered numbers (prereg §11, G.3).")
    say("=" * 78)
    return {"plan": len(plan), "cached": cached, "todo": len(todo), "per_model": est_rows,
            "est_usd": total}


# ───────────────────────────────────────────────────────────────── the HTTP call
def api_key(env: str = "GATEWAY_KEY") -> str:
    k = os.environ.get(env)
    if not k:
        die(f"{env} environment variable is not set — the key is never read from a file")
    return k


def http_chat(cfg: Config, body: dict, timeout: int, role: str = "subject") -> tuple[int, dict, str]:
    if not LIVE_ALLOWED:                    # belt and braces: unreachable in dry runs
        raise RuntimeError("live call attempted while LIVE_ALLOWED is False")
    base, key_env = cfg.route(role)
    data = json.dumps(body, ensure_ascii=False).encode("utf-8")
    req = urlrequest.Request(
        base + "/chat/completions", data=data,
        headers={"Authorization": "Bearer " + api_key(key_env), "Content-Type": "application/json"})
    try:
        with urlrequest.urlopen(req, timeout=timeout) as r:
            raw = r.read().decode("utf-8", "replace")
            return r.status, json.loads(raw), ""
    except urlerror.HTTPError as e:
        raw = e.read().decode("utf-8", "replace")
        try:
            return e.code, json.loads(raw), ""
        except json.JSONDecodeError:
            return e.code, {}, raw[:500]
    except Exception as e:                  # timeouts, DNS, connection resets
        return 0, {}, f"{type(e).__name__}: {e}"[:500]


def extract_meta(resp: dict) -> dict:
    """finish_reason, full usage, the provider's model string, any fingerprint field (C.2, C.7)."""
    ch = (resp.get("choices") or [{}])[0]
    msg = ch.get("message") or {}
    fp = None
    for k, v in resp.items():
        if "fingerprint" in k.lower() and isinstance(v, (str, int)):
            fp = f"{k}={v}"
            break
    return {
        "finish_reason": ch.get("finish_reason"),
        "usage": resp.get("usage") or {},
        "model_string": resp.get("model"),
        "fingerprint": fp,
        "content": msg.get("content") or "",
        "content_chars": len(msg.get("content") or ""),
        "response_id": resp.get("id"),
    }


def update_window(logs: Path, phase: str, rec: dict):
    """Run-window log for C.7 / H.4: first and last call, observed model strings."""
    p = logs / "run_window.json"
    with _WINDOW_LOCK:
        w = json.loads(p.read_text(encoding="utf-8")) if p.exists() else {"phases": {}, "models": {}}
        ph = w["phases"].setdefault(phase, {"first_call": None, "last_call": None, "calls": 0})
        ts = rec["t_request"]
        ph["first_call"] = min(ph["first_call"] or ts, ts)
        ph["last_call"] = max(ph["last_call"] or ts, ts)
        ph["calls"] += 1
        m = w["models"].setdefault(rec["spec"]["model_id"], {"model_strings": {}, "fingerprints": {}})
        ms = rec["meta"].get("model_string")
        if ms:
            e = m["model_strings"].setdefault(str(ms), {"first_seen": ts, "count": 0})
            e["count"] += 1
            e["last_seen"] = ts
        fp = rec["meta"].get("fingerprint")
        if fp:
            e = m["fingerprints"].setdefault(str(fp), {"first_seen": ts, "count": 0})
            e["count"] += 1
            e["last_seen"] = ts
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(json.dumps(w, indent=1, ensure_ascii=False), encoding="utf-8")


def build_body(cfg: Config, spec: dict, params: dict) -> dict:
    """The ONLY place a request body is built (M4). Bare single-turn prompt: exactly one user
    message, no system role, no history, no tools (B.3). Extractors add response_format only."""
    body = {"model": spec["model_id"],
            "messages": [{"role": "user", "content": spec["prompt"]}],
            **params}
    if spec["role"] == "extractor":
        if "response_format" in params:
            rf = body.pop("response_format")          # per-role override (v0.54)
            if rf is not None:                        # None = deliberately omitted
                body["response_format"] = rf if isinstance(rf, dict) else {"type": "json_object"}
        elif cfg.extractor_defaults.get("response_format") == "json":
            body["response_format"] = {"type": "json_object"}
    return body


def do_one_call(cfg: Config, cache: Cache, logs: Path, spec: dict, params: dict,
                timeout: int, retries: int) -> str:
    """Returns 'cached' | 'ok' | 'error'. Never raises for an API failure."""
    k = Cache.key(spec, params)
    if cache.has(spec, k):
        return "cached"
    body = build_body(cfg, spec, params)
    attempts, status, resp, err, ok = [], 0, {}, "", False
    t0 = now_iso()
    start = time.time()
    for n in range(1, retries + 2):
        status, resp, err = http_chat(cfg, body, timeout, spec.get("role", "subject"))
        ok = status == 200 and (resp.get("choices") or [{}])[0].get("message") is not None
        wait = 0.0
        if not ok and n <= retries and (status in (0, 408, 409, 425, 429, 500, 502, 503, 504)):
            wait = min(60.0, 2.0 ** n) * (0.6 + 0.8 * random.random())   # backoff + jitter
        attempts.append({"n": n, "http_status": status, "error": err or None,
                         "waited_s": round(wait, 2)})
        if ok:
            break
        if wait:
            time.sleep(wait)
        else:
            break
    rec = {"schema": SCHEMA, "key": k, "spec": spec, "params": params,
           "request": {"url": cfg.route(spec.get("role", "subject"))[0] + "/chat/completions",
                       "body": body},
           "t_request": t0, "t_response": now_iso(), "duration_s": round(time.time() - start, 3),
           "http_status": status, "attempts": attempts, "error": err or None,
           "response": resp, "meta": {}, "models_json_sha256": cfg.sha256}
    # B2: `ok` already requires choices[0].message; a gateway error payload delivered with
    # HTTP 200 ({"error": {...}}) must never be cached as a success — it would be coded
    # degenerate, burn both redraws and report a cell short although the model never answered.
    good = ok and not resp.get("error")
    if good:
        rec["meta"] = extract_meta(resp)
    else:
        rec["error"] = err or (json.dumps(resp.get("error"), ensure_ascii=False)[:500]
                               if resp.get("error") else "no choices[0].message in response")
    cache.write(spec, k, rec, err=not good)
    if good:
        update_window(logs, spec["phase"], rec)
        return "ok"
    append_jsonl(logs / "failed.jsonl",
                 {"ts": now_iso(), "key": k, "spec": spec, "http_status": status,
                  "error": (err or json.dumps(resp)[:300])})
    return "error"


def execute(cfg: Config, cache: Cache, logs: Path, plan: list[dict], params_for,
            concurrency: int, timeout: int, retries: int, label: str):
    done = Counter()
    t0 = time.time()
    def work(s):
        return do_one_call(cfg, cache, logs, s, params_for(s), timeout, retries)
    with ThreadPoolExecutor(max_workers=concurrency) as ex:
        futs = {ex.submit(work, s): s for s in plan}
        for i, f in enumerate(as_completed(futs), 1):
            done[f.result()] += 1
            if i % 25 == 0 or i == len(plan):
                el = time.time() - t0
                rate = i / el if el else 0
                say(f"  [{label}] {i}/{len(plan)}  ok={done['ok']} cached={done['cached']} "
                    f"err={done['error']}  {rate:.1f} calls/s  eta {(len(plan)-i)/rate/60:.1f} min"
                    if rate else f"  [{label}] {i}/{len(plan)}")
    say(f"  [{label}] finished: {dict(done)}")
    return done


# ───────────────────────────────────────────────────── preflight (the --go gate)
def validator_clean(returncode: int, stdout: str) -> bool:
    """B1: the validator exits 1 on any FAIL; the anchored regex is a second, independent read.
    A substring test would accept '10 FAIL' because it contains '0 FAIL'."""
    return returncode == 0 and re.search(r"^0 FAIL\b", stdout, re.M) is not None


def preflight(a, cfg: Config, queries: list[dict], plan: list[dict], mode: str) -> None:
    """Everything that must be true before ONE cent is spent. Failing any item aborts."""
    global LIVE_ALLOWED
    checks = []

    def chk(ok, label, detail=""):
        checks.append((bool(ok), label, detail))

    # Preflight the route this phase will actually use. With subjects and the extractor on
    # different gateways, checking only GATEWAY_KEY would pass and then fail on the first call.
    role = "extractor" if mode == "extract" else "subject"
    r_base, r_env = cfg.route(role)
    chk(os.environ.get(r_env), f"{r_env} is set in the environment  (route for {role} calls)")
    if cfg.extractor_base != cfg.base:
        chk(True, f"split routing: subjects -> {cfg.base} | extractor -> {cfg.extractor_base}",
            "both are recorded in models.json and hashed at freeze (H.4); every extraction in a"
            " phase must use ONE route")
    qp = Path(a.queries)
    chk(qp.exists(), f"{qp} exists")
    shafile = Path(str(qp) + ".sha256")
    if shafile.exists() and qp.exists():
        recorded = shafile.read_text(encoding="utf-8").split()[0]
        chk(recorded == sha256_file(qp), "queries.csv matches its recorded sha256",
            "re-run `export` if the kit changed, and re-record the hash in prereg H.4")
    else:
        chk(False, "queries.csv.sha256 present (H.4 hash record)",
            "run `export` to create it")

    if mode == "smoke":
        smoke = [q for q in queries if (q.get("smoke_test") or "").upper() == "YES"]
        chk(len(smoke) == 10, f"exactly 10 smoke_test rows exported and FINAL (found {len(smoke)})",
            "export only writes FINAL rows — write and approve the 10 smoke rows first")
        full = [q for q in smoke if all((q.get(ARM_COLUMN[x]) or "").strip() for x in ARMS)]
        chk(len(full) == len(smoke) and smoke, "every smoke row has all three renderings")
        cats = {q.get("category") for q in smoke}
        chk(len(cats) == 10 if smoke else False, "smoke rows cover 10 distinct categories")

    v = Path(a.validator)
    if v.exists() and Path(a.kit).exists():
        import subprocess
        r = subprocess.run([sys.executable, str(v), str(a.kit)], capture_output=True, text=True)
        tail = (r.stdout or "").strip().splitlines()[-1:] or [""]
        chk(validator_clean(r.returncode, r.stdout or ""), "validate_queries.py exits 0 with 0 FAIL", tail[0])
    else:
        chk(False, "validate_queries.py runnable against the kit", f"{v} or {a.kit} missing")

    chk(plan, "the plan contains at least one call to make")
    chk(len(plan) <= a.max_calls, f"plan size {len(plan)} is within --max-calls {a.max_calls}",
        "raise --max-calls deliberately if this is really the run you intend")

    say("")
    say("PREFLIGHT")
    for ok, label, detail in checks:
        say(f"  [{'PASS' if ok else 'FAIL'}] {label}" + (f"   → {detail}" if detail and not ok else ""))
    bad = [c for c in checks if not c[0]]
    if bad:
        die(f"{len(bad)} preflight check(s) failed — no call was made")
    if not a.go:
        die("preflight passed, but --go was not given. This is the spend gate: re-run with --go.")
    LIVE_ALLOWED = True
    say("\n  --go accepted. Live calls enabled for this process only.\n")


# ────────────────────────────────────────────────────────────────── run command
def cmd_run(a):
    cfg = Config(Path(a.models))
    queries = load_queries(Path(a.queries))
    cache = Cache(Path(a.runs))
    logs = Path(a.runs) / "logs"
    prices = load_prices(Path(a.prices))

    if a.smoke:
        mode, phase, reps = "smoke", "smoke", a.reps or 2
        rows = [q for q in queries if (q.get("smoke_test") or "").upper() == "YES"]
        arms = ARMS + ((ARM_TRANSLIT,) if not a.no_translit else ())
        live_arms = [x for x in arms if any((q.get(ARM_COLUMN[x]) or "").strip() for q in rows)]
        title = (f"SMOKE — {len(rows)} smoke_test rows × {len(live_arms)} arms "
                 f"({'+'.join(live_arms)}) × {len(cfg.subjects)} subjects × {reps} reps")
        if ARM_TRANSLIT in arms and ARM_TRANSLIT not in live_arms:
            title += "  [no bl_translit text in these rows -> translit arm empty]"
    elif a.pilot:
        mode, phase, reps = "pilot", "pilot", a.reps or 3
        rows = [q for q in queries if (q.get("calibration") or "").upper() == "YES"]
        arms = ARMS
        title = f"PILOT — {len(rows)} calibration rows × 3 arms × {len(cfg.subjects)} subjects × {reps} reps"
    elif a.translit:
        if not a.reps:
            die("--translit needs an explicit --reps equal to the calibrated r from the pilot "
                "(G.1: 5 or 7). No default on purpose (M6).")
        mode, phase, reps = "translit", "main", a.reps
        rows = [q for q in queries if (q.get("robust50") or "").upper() == "YES"]
        arms = (ARM_TRANSLIT,)
        title = f"TRANSLIT ARM — {len(rows)} robust50 rows × bl_translit × {len(cfg.subjects)} subjects × {reps} reps"
    elif a.main:
        mode, phase, reps = "main", "main", a.reps or 5
        rows = queries
        arms = ARMS
        title = f"MAIN RUN — {len(rows)} queries × 3 arms × {len(cfg.subjects)} subjects × {reps} reps"
    else:
        die("choose one of --smoke / --pilot / --main / --translit")

    if not rows:
        die(f"no rows selected for {mode} — are the relevant rows FINAL and exported?")
    plan = build_subject_plan(cfg, rows, arms, reps, phase, a.seed)
    means = measured_means(cache, ["smoke", "pilot", "main"])
    todo = [s for s in plan
            if not cache.has(s, Cache.key(s, cfg.subject_params(
                next(x for x in cfg.subjects if x["id"] == s["model_id"]))))]
    print_plan(cfg, plan, cache, prices, means, title)

    if a.dry_run or not a.go:
        say("DRY RUN — nothing was called. Add --go (after preflight) to execute.\n")
        if not a.dry_run:
            preflight(a, cfg, queries, todo, mode)      # shows the gate even without --go
        return
    preflight(a, cfg, queries, todo, mode)

    params_for = lambda s: cfg.subject_params(
        next(x for x in cfg.subjects if x["id"] == s["model_id"]))
    append_jsonl(logs / "runs.jsonl", {"ts": now_iso(), "mode": mode, "phase": phase,
                                       "arms": list(arms), "reps": reps, "rows": len(rows),
                                       "planned": len(plan), "to_call": len(todo),
                                       "seed": a.seed, "models_json_sha256": cfg.sha256})
    execute(cfg, cache, logs, todo, params_for, a.concurrency, a.timeout, a.retries, mode)
    say("\nnext: `code --phase %s` to assign outcome codes and build the redraw plan.\n" % phase)


# ────────────────────────────────────────────── deterministic outcome coding
BANGLA_RE = re.compile(r"[ঀ-৿]")
LATIN_RE = re.compile(r"[A-Za-z]")


def bangla_ratio(text: str) -> float:
    b = len(BANGLA_RE.findall(text))
    l = len(LATIN_RE.findall(text))
    return (b / (b + l)) if (b + l) else 0.0


def bangla_ratio_words(text: str) -> float:
    toks = [t for t in re.split(r"\s+", text or "") if BANGLA_RE.search(t) or LATIN_RE.search(t)]
    if not toks:
        return 0.0
    bn = sum(1 for t in toks if len(BANGLA_RE.findall(t)) >= len(LATIN_RE.findall(t)))
    return bn / len(toks)


PAREN_RE = re.compile(r"\([^)]*\)|\[[^\]]*\]|（[^）]*）")
DIGIT_RE = re.compile(r"\d")            # Unicode-aware: Western 0-9 and Bangla ০-৯ alike
LEAD_PUNCT = "\"'“‘(«[*-–—•·"


def prose_ratio(text: str) -> float:
    """D.4 (approved 2 Sep 2026): measure PROSE, not characters. Strip parenthesised chunks,
    drop tokens containing digits, drop Latin-only tokens whose first letter is uppercase
    (product names, acronyms: Galaxy, Redmi, Primo, GB, RAM), then Bangla letters /
    (Bangla + Latin letters) on what remains. No alias-table dependency. If nothing remains,
    fall back to the plain character ratio."""
    t = PAREN_RE.sub(" ", text or "")
    kept = []
    for tok in re.split(r"\s+", t):
        if not tok or DIGIT_RE.search(tok):
            continue
        core = tok.lstrip(LEAD_PUNCT)
        if core and not BANGLA_RE.search(tok) and LATIN_RE.search(tok) and core[0].isupper():
            continue
        kept.append(tok)
    rest = " ".join(kept)
    b, l = len(BANGLA_RE.findall(rest)), len(LATIN_RE.findall(rest))
    if b + l == 0:
        return bangla_ratio(text)
    return b / (b + l)


def script_class(text: str) -> str:
    r = prose_ratio(text)
    if r >= BANGLA_RATIO_BN:
        return "bn"
    if r <= BANGLA_RATIO_LATIN:
        return "latin"
    return "mixed"


def degenerate_check(content: str, finish_reason: str) -> tuple[bool, str]:
    """Empty / loop / unreadable = degenerate (D.5 code 4). Truncation is NOT degenerate (D.6)."""
    s = (content or "").strip()
    if len(s) < DEGEN_MIN_CHARS:
        return True, "empty"
    bad = sum(1 for ch in s if ch == "�" or (unicodedata.category(ch) == "Cc" and ch not in "\n\r\t"))
    if bad / max(1, len(s)) > DEGEN_REPLACEMENT_RATIO:
        return True, "unreadable_chars"
    # tail-anchored periodicity test: a degenerating model repeats a unit until the token cap.
    # For each candidate period p, the last span characters must satisfy t[i] == t[i-p]
    # (phase-independent, so it catches loops whatever offset they start at).
    t = re.sub(r"\s+", " ", s).strip()
    n = len(t)
    for period in range(1, DEGEN_LOOP_MAX_PERIOD + 1):
        reps = max(DEGEN_LOOP_REPEATS, -(-DEGEN_LOOP_MIN_SPAN // period))       # ceil division
        span = period * reps
        if n < span:
            continue
        tail = t[-span:]
        if tail[:period].strip() and all(tail[i] == tail[i - period] for i in range(period, span)):
            return True, f"loop:period{period}x{reps}"
    return False, ""


class RefusalDetector:
    """F4: refusal phrase patterns authored for this study in English, Bangla script and Banglish, one regex per line, hashed at freeze; a match codes a refusal only under the D.5 content conjunct. Loaded from refusal_rules.md; the conjunct is applied in `finalize`.
    Format: one regex per line, compiled with re.IGNORECASE; '#' comments and blank lines ignored.
    Without the file the coder does NOT guess — `refused` stays null and the run is marked
    incomplete, unless --allow-missing-refusal-rules is passed for structural testing."""

    def __init__(self, path: Path, allow_missing: bool):
        self.path, self.rules, self.available = path, [], False
        if path.exists():
            for i, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
                line = line.strip()
                if not line or line.startswith("#"):
                    continue
                try:
                    self.rules.append((f"{path.name}:{i}", re.compile(line, re.I)))
                except re.error as e:
                    die(f"{path}:{i} is not a valid regex: {e}")
            self.available = True
            self.sha256 = sha256_file(path)
        elif not allow_missing:
            die(f"{path} not found (prereg F4). Import the thesis refusal rules first, or pass "
                f"--allow-missing-refusal-rules to code structure only (refused will be null).")

    def check(self, content: str):
        if not self.available:
            return None, None
        for name, rx in self.rules:
            if rx.search(content or ""):
                return True, name
        return False, None


def conformity(arm: str, sclass: str):
    """-> (answer_language | None, pending_extractor, conforming | None)   [code stage]

    D.4: the deterministic script check owns the bn-vs-Latin boundary. Everything the extractor
    must decide (en vs banglish inside Latin text; the language of a `mixed` answer, M1) is
    left pending here and resolved by `finalize`.
      bn arm          : 'bn' -> valid; 'latin' -> language_reversion (decided here); 'mixed' -> pending
      en/bl/translit  : 'bn' -> language_reversion (decided here); 'latin'/'mixed' -> pending
    D.4 defines *reversion* for the bn and bl arms; a non-conforming en-arm answer is coded the
    same way so the query-language x answer-language matrix is complete, but it is matrix-only,
    not an RQ5 confirmatory contrast (approved 2 Sep 2026)."""
    if sclass == "bn":
        return "bn", False, (arm == "bn")
    if sclass == "latin" and arm == "bn":
        return None, False, False
    return None, True, None


ARM_TARGET = {"bn": "bn", "en": "en", "bl": "banglish", ARM_TRANSLIT: "banglish"}


def resolve_language(arm: str, sclass: str, ext_lang):
    """-> (answer_language, conforming | None, adjudicate_reason | None)   [finalize stage]
    The deterministic verdict keeps authority on the bn/Latin boundary; disagreements with the
    extractor are listed for R1 (D.4). `mixed` is decided by the extractor (M1)."""
    target = ARM_TARGET[arm]
    if sclass == "bn":
        reason = None if ext_lang in (None, "bn") else f"script=bn but extractor={ext_lang}"
        return "bn", target == "bn", reason
    if sclass == "latin":
        if ext_lang == "bn":
            reason = "script=latin but extractor=bn"
            return "latin_unresolved", (False if arm == "bn" else None), reason
        if ext_lang is None:
            return None, (False if arm == "bn" else None), None
        return ext_lang, (False if arm == "bn" else ext_lang == target), None
    # mixed: extractor decides (M1)
    if ext_lang is None:
        return "mixed", None, None
    return ext_lang, ext_lang == target, None


def _pct(vals, q):
    if not vals:
        return None
    v = sorted(vals)
    i = min(len(v) - 1, max(0, int(round(q * (len(v) - 1)))))
    return v[i]


def print_outcome_tables(rows: list[dict], stage: str):
    """Shared summary for `code` (provisional) and `finalize` (final)."""
    c = Counter(r["outcome"] for r in rows)
    say(f"\n{stage.upper()} OUTCOMES — {len(rows)} answers")
    for k, v in c.most_common():
        say(f"  {k:<24}{v:>7}  ({v/max(1,len(rows)):.1%})")
    hedged = sum(1 for r in rows if r.get("hedged"))
    if hedged:
        say(f"  hedged (F4 fired, content kept){hedged:>7}  reported descriptively (D.5 refinement)")
    say(f"\noutcome codes by arm ({stage}):")
    arms_seen = sorted({r["arm"] for r in rows})
    bad_keys = ("refusal", "degenerate") if stage == "final" else ("refusal_candidate", "degenerate")
    for arm in arms_seen:
        sub = [r for r in rows if r["arm"] == arm]
        cc = Counter(r["outcome"] for r in sub)
        bad = sum(cc[k] for k in bad_keys) / max(1, len(sub))
        note = "   <- H.2: reported separately; >70% per model => F.10 descriptive" if arm == ARM_TRANSLIT else ""
        say(f"  {arm:<12}n={len(sub):<6}" + " ".join(f"{k}={v}" for k, v in cc.most_common())
            + f"   {'+'.join(bad_keys)}={bad:.1%}{note}")
    if ARM_TRANSLIT in arms_seen:
        say(f"\n  per-model {'+'.join(bad_keys)} on the bl_translit arm (H.2 gate, C.6 threshold 70%):")
        for mid in sorted({r["model_id"] for r in rows if r["arm"] == ARM_TRANSLIT}):
            sub = [r for r in rows if r["arm"] == ARM_TRANSLIT and r["model_id"] == mid]
            cc = Counter(r["outcome"] for r in sub)
            bad = sum(cc[k] for k in bad_keys) / max(1, len(sub))
            flag = "  ** EXCEEDS GATE -> F.10 descriptive for this model" if bad > 0.70 else ""
            say(f"    {mid:<24}{bad:>6.1%}  (n={len(sub)}){flag}")
    # M3: reasoning burn — evidence for keeping or raising max_tokens before it is frozen (C.2)
    say("\n  finish_reason=length by model (M3 — max_tokens evidence, C.2):")
    say(f"    {'model':<24}{'answers':>8}{'length':>8}{'empty+length':>14}{'share length':>14}")
    for mid in sorted({r["model_id"] for r in rows}):
        sub = [r for r in rows if r["model_id"] == mid]
        ln = sum(1 for r in sub if r.get("truncated"))
        el = sum(1 for r in sub if r.get("truncated") and r.get("degenerate_reason") == "empty")
        say(f"    {mid:<24}{len(sub):>8}{ln:>8}{el:>14}{ln/max(1,len(sub)):>14.1%}")
    # M1: bn-arm bangla_ratio distribution — the data on which the D.4 thresholds get approved
    bn_rows = [r for r in rows if r["arm"] == "bn" and not r.get("degenerate")]
    if bn_rows:
        say(f"\n  bn-arm script ratios over {len(bn_rows)} non-degenerate answers (D.4 — cuts "
            f"{BANGLA_RATIO_LATIN}/{BANGLA_RATIO_BN} are [PROPOSED]; approve on this, then validate "
            f"against the 300 labels stratified by script_class):")
        for label, key, used in (("prose", "prose_ratio", True), ("chars", "bangla_ratio", False),
                                 ("words", "bangla_ratio_words", False)):
            v = [r.get(key) for r in bn_rows if r.get(key) is not None]
            if not v:
                continue
            below = sum(1 for x in v if x < BANGLA_RATIO_BN)
            say(f"    {label:<6}: min {min(v):.3f}   p10 {_pct(v, .10):.3f}   median {_pct(v, .5):.3f}   "
                f"p90 {_pct(v, .90):.3f}   max {max(v):.3f}   below {BANGLA_RATIO_BN}: {below} ({below/len(v):.1%})"
                + ("   <- decides script_class" if used else "   (information only)"))
        say("    script_class: " + ", ".join(f"{k}={v}" for k, v in
                                             Counter(r["script_class"] for r in bn_rows).items()))


def cmd_code(a):
    cache = Cache(Path(a.runs))
    det = RefusalDetector(Path(a.refusal_rules), a.allow_missing_refusal_rules)
    out_path = Path(a.runs) / "coded" / f"{a.phase}.jsonl"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    rows, per_cell = [], defaultdict(list)
    for rec in cache.records(a.phase):
        meta, spec = rec.get("meta") or {}, rec["spec"]
        content = meta.get("content") or ""
        fr = meta.get("finish_reason")
        degen, why = degenerate_check(content, fr)
        f4, rule = (False, None) if degen else det.check(content)
        sclass = script_class(content)
        lang, pending, conforming = conformity(spec["arm"], sclass)
        if degen:
            outcome = "degenerate"
        elif f4 is None:
            outcome = "pending_refusal_rules"
        elif f4:
            outcome = "refusal_candidate"          # B4: conjunct with extraction, resolved in finalize
        elif pending:
            outcome = "pending_extractor"
        else:
            outcome = "valid" if conforming else "language_reversion"
        row = {"key": rec["key"], "phase": spec["phase"], "model_id": spec["model_id"],
               "query_id": spec["query_id"], "arm": spec["arm"], "rep": spec["rep"],
               "draw": spec["draw"], "finish_reason": fr, "truncated": fr == "length",
               "content_chars": meta.get("content_chars"), "bangla_ratio": round(bangla_ratio(content), 4),
               "bangla_ratio_words": round(bangla_ratio_words(content), 4),
               "prose_ratio": round(prose_ratio(content), 4),
               "script_class": sclass, "f4_fires": f4, "f4_rule": rule,
               "degenerate": degen, "degenerate_reason": why or None,
               "answer_language": lang, "pending_extractor": pending, "outcome": outcome,
               "model_string": meta.get("model_string"), "t_request": rec.get("t_request")}
        rows.append(row)
        per_cell[(spec["model_id"], spec["query_id"], spec["arm"], spec["rep"])].append(row)
    with open(out_path, "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    # redraw plan: degenerate only, at most MAX_DRAWS total draws, content never repaired (D.6)
    redraws, short = [], []
    for (mid, qid, arm, rep), rs in sorted(per_cell.items()):
        rs.sort(key=lambda x: x["draw"])
        if any(not r["degenerate"] for r in rs):
            continue                                  # cell already has a usable answer
        last = rs[-1]
        if last["draw"] < MAX_DRAWS:
            redraws.append({"model_id": mid, "query_id": qid, "arm": arm, "rep": rep,
                            "from_draw": last["draw"], "to_draw": last["draw"] + 1,
                            "reason": last["degenerate_reason"]})
        else:
            short.append({"model_id": mid, "query_id": qid, "arm": arm, "rep": rep,
                          "draws_used": last["draw"], "reason": last["degenerate_reason"]})
    rp = Path(a.runs) / "coded" / f"{a.phase}.redraw_plan.json"
    rp.write_text(json.dumps({"generated": now_iso(), "phase": a.phase, "redraws": redraws,
                              "persistently_degenerate": short}, indent=1), encoding="utf-8")

    say(f"\ncoded {len(rows)} answers from phase '{a.phase}' -> {out_path}")
    print_outcome_tables(rows, "provisional")
    if redraws:
        say(f"\n{len(redraws)} degenerate cell(s) eligible for redraw -> {rp}")
        say("   run:  n1_pipeline.py redraw --phase %s --go" % a.phase)
    if short:
        say(f"{len(short)} cell(s) persistently degenerate after {MAX_DRAWS} draws — left short and reported (D.6)")
    if not det.available:
        say("\nWARNING: no refusal rules loaded (F4) — f4_fires is null and outcomes are "
            "incomplete. This coding is structural only and must be re-run before analysis.")
    say("\nnext: extract --phase %s --primary --go, then finalize --phase %s" % (a.phase, a.phase))


def cmd_finalize(a):
    """B3 + B4: fold the primary extraction back into the outcome codes (offline).
    - en-vs-banglish (and the language of `mixed` answers) comes from the extractor;
      the deterministic verdict keeps the bn/Latin boundary, disagreements go to R1 (D.4)
    - refusal is conjunctive (D.5 code 3, approved 2 Sep 2026): F4 fires AND the extraction has
      no brands, no retailers and no prices; F4 fires with content present -> content code kept,
      hedged = true"""
    cache = Cache(Path(a.runs))
    coded_p = Path(a.runs) / "coded" / f"{a.phase}.jsonl"
    coded = read_jsonl(coded_p)
    if not coded:
        die(f"{coded_p} missing or empty — run `code --phase {a.phase}` first")
    ext_p = Path(a.runs) / "extracted" / f"{a.phase}.primary.jsonl"
    ext_rows = read_jsonl(ext_p)
    if not ext_rows:
        die(f"{ext_p} missing or empty — run `extract --phase {a.phase} --primary --go` first")
    ext = {r["subject_key"]: r for r in ext_rows}
    final, adjud = [], []
    for row in coded:
        e = ext.get(row["key"])
        extraction = e["extraction"] if (e and e.get("extraction") is not None
                                         and not e.get("schema_errors")) else None
        ext_lang = extraction.get("answer_language") if extraction else None
        has_content = (bool(extraction.get("brands") or extraction.get("retailers")
                            or extraction.get("prices")) if extraction else None)   # D.5 code 3 conjunct
        out = dict(row, extractor_answer_language=ext_lang,
                   extractor_refused=(extraction.get("refused") if extraction else None),
                   extractor_has_content=has_content, hedged=False, adjudicate=False,
                   adjudicate_reason=None, extraction_status=("ok" if extraction else
                                                              ("schema_fail" if e else "missing")))
        arm, sclass, f4 = row["arm"], row["script_class"], row.get("f4_fires")
        if row["degenerate"]:
            out["outcome"] = "degenerate"
        elif f4 is None:
            out["outcome"] = "pending_refusal_rules"
        elif extraction is None and (f4 or row["pending_extractor"]):
            out["outcome"] = "pending_extractor"       # extraction missing or schema-invalid
        elif f4 and not has_content:
            out["outcome"] = "refusal"                 # code 3: F4 fires AND no brands, retailers or prices
        else:
            if f4 and has_content:
                out["hedged"] = True                   # detector fired, content present: keep code
            lang, conforming, reason = resolve_language(arm, sclass, ext_lang)
            out["answer_language"] = lang
            if reason:
                out["adjudicate"], out["adjudicate_reason"] = True, reason
            out["outcome"] = ("pending_adjudication" if conforming is None
                              else "valid" if conforming else "language_reversion")
        final.append(out)
        if out["adjudicate"]:
            snippet = ""
            try:
                recp = cache.path({"phase": row["phase"], "model_id": row["model_id"],
                                   "query_id": row["query_id"], "arm": arm, "rep": row["rep"],
                                   "draw": row["draw"]}, row["key"])
                snippet = ((json.loads(recp.read_text(encoding="utf-8")).get("meta") or {})
                           .get("content") or "")[:300]
            except Exception:
                pass
            adjud.append({"key": row["key"], "query_id": row["query_id"], "arm": arm,
                          "model_id": row["model_id"], "rep": row["rep"], "draw": row["draw"],
                          "script_class": sclass, "bangla_ratio": row["bangla_ratio"],
                          "extractor_answer_language": ext_lang, "reason": out["adjudicate_reason"],
                          "provisional_outcome": out["outcome"], "content_start": snippet,
                          "R1_decision": None, "R1_note": None})
    fp = coded_p.with_name(f"{a.phase}.final.jsonl")
    with open(fp, "w", encoding="utf-8") as f:
        for r in final:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    ap_ = coded_p.with_name(f"{a.phase}.adjudicate.jsonl")
    with open(ap_, "w", encoding="utf-8") as f:
        for r in adjud:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    say(f"\nfinalized {len(final)} answers -> {fp}")
    st = Counter(r["extraction_status"] for r in final)
    say(f"extraction status: " + ", ".join(f"{k}={v}" for k, v in st.items()))
    print_outcome_tables(final, "final")
    say(f"\n{len(adjud)} deterministic-vs-extractor disagreement(s) on the bn/Latin boundary "
        f"-> {ap_}  (for R1, blind to model; D.4)")
    pend = sum(1 for r in final if r["outcome"].startswith("pending"))
    if pend:
        say(f"WARNING: {pend} answer(s) still pending — analysis must not start until they are 0.")


def parse_extractor_json(text: str):
    """M2: tolerant parse of the EXTRACTOR's output (not a subject answer — no D.6 conflict):
    strip ``` fences, then fall back to the first balanced {...} block."""
    t = (text or "").strip()
    m = re.match(r"^```[a-zA-Z]*\s*(.*?)\s*```$", t, re.S)
    if m:
        t = m.group(1).strip()
    try:
        return json.loads(t), None
    except json.JSONDecodeError:
        pass
    i = t.find("{")
    if i < 0:
        return None, "not JSON: no object found"
    depth = 0
    for j in range(i, len(t)):
        if t[j] == "{":
            depth += 1
        elif t[j] == "}":
            depth -= 1
            if depth == 0:
                try:
                    return json.loads(t[i:j + 1]), None
                except json.JSONDecodeError as e:
                    return None, f"not JSON: {e}"
    return None, "not JSON: unbalanced braces"


def cmd_redraw(a):
    cfg = Config(Path(a.models))
    queries = {q["query_id"]: q for q in load_queries(Path(a.queries))}
    cache = Cache(Path(a.runs))
    logs = Path(a.runs) / "logs"
    rp = Path(a.runs) / "coded" / f"{a.phase}.redraw_plan.json"
    if not rp.exists():
        die(f"{rp} not found — run `code --phase {a.phase}` first")
    plan_j = json.loads(rp.read_text(encoding="utf-8"))
    specs = []
    for r in plan_j["redraws"]:
        q = queries.get(r["query_id"])
        if not q:
            continue
        text = (q.get(ARM_COLUMN[r["arm"]]) or "").strip()
        subj = next((x for x in cfg.subjects if x["id"] == r["model_id"]), None)
        if not text or not subj:
            continue
        specs.append({"phase": a.phase, "role": "subject", "model_id": subj["id"],
                      "model_name": subj.get("name", subj["id"]), "query_id": r["query_id"],
                      "category": q.get("category", ""), "arm": r["arm"], "rep": r["rep"],
                      "draw": r["to_draw"], "prompt": text, "prompt_sha256": sha256_text(text)})
    say(f"redraw plan: {len(specs)} call(s) at shifted draw ids (max {MAX_DRAWS} draws per cell)")
    if not specs:
        return
    if a.dry_run or not a.go:
        for s in specs[:20]:
            say(f"  {s['model_id']:<22}{s['query_id']} {s['arm']} rep{s['rep']} draw{s['draw']}")
        say("DRY RUN — nothing was called. Add --go to execute.")
        return
    preflight(a, cfg, list(queries.values()), specs, "redraw")
    for r in plan_j["redraws"]:
        append_jsonl(logs / "redraws.jsonl", {"ts": now_iso(), "phase": a.phase, **r,
                                              "rule": f"D.6 redraw {r['to_draw']}/{MAX_DRAWS}, content never repaired"})
    params_for = lambda s: cfg.subject_params(
        next(x for x in cfg.subjects if x["id"] == s["model_id"]))
    execute(cfg, cache, logs, specs, params_for, a.concurrency, a.timeout, a.retries, "redraw")
    say("re-run `code` to fold the redraws in.")


# ───────────────────────────────────────────────────────── extraction (E.1, E.2)
EXTRACTION_SCHEMA_DOC = """{
  "brands":      ["<canonical brand name>", ...],
  "recommended": ["<subset of brands the answer actually advises buying>", ...],
  "prices":      [{"amount": <number>, "currency": "BDT" | "USD" | "other" | "unstated"}, ...],
  "retailers": ["<shop or marketplace name>", ...],
  "refused":   true | false,
  "answer_language": "en" | "bn" | "banglish" | "mixed" | "other"
}"""
CURRENCIES = {"BDT", "USD", "other", "unstated"}
LANGS = {"en", "bn", "banglish", "mixed", "other"}


def validate_extraction(obj) -> list[str]:
    """Hand-rolled JSON-schema check (stdlib only) — E.1 requires schema-valid extractor output."""
    e = []
    if not isinstance(obj, dict):
        return ["top level is not an object"]
    for k in ("brands", "retailers"):
        v = obj.get(k)
        if not isinstance(v, list):
            e.append(f"{k} must be a list")
        elif any(not isinstance(x, str) for x in v):
            e.append(f"{k} must contain only strings")
    p = obj.get("prices")
    if not isinstance(p, list):
        e.append("prices must be a list")
    else:
        for i, it in enumerate(p):
            if not isinstance(it, dict):
                e.append(f"prices[{i}] must be an object")
                continue
            if not isinstance(it.get("amount"), (int, float)):
                e.append(f"prices[{i}].amount must be a number")
            if it.get("currency") not in CURRENCIES:
                e.append(f"prices[{i}].currency must be one of {sorted(CURRENCIES)}")
    if "recommended" in obj:                       # optional; RQ2 sensitivity (D.2)
        rc = obj["recommended"]
        if not isinstance(rc, list) or any(not isinstance(x, str) for x in rc):
            e.append("recommended must be a list of strings")
        elif isinstance(obj.get("brands"), list):
            # Brand identity is CASE-INSENSITIVE here (registered 11 Sep 2026, prereg v0.35): the
            # instruction returns non-table brands as surface forms, so one brand can legally appear
            # as 'Walton' in one field and 'walton' in the other, and both smoke-test failures under
            # the strict check were exactly that. The check still catches what it exists to catch —
            # a recommendation naming a brand ABSENT from `brands`. Accent differences still fail:
            # 'Estée' vs 'Estee' is a real inconsistency for the alias table (F2), not letter case.
            have = {x.casefold() for x in obj["brands"] if isinstance(x, str)}
            extra = [x for x in rc if x.casefold() not in have]
            if extra:
                e.append(f"recommended must be a subset of brands, case-insensitively"
                         f" (not in brands: {extra[:3]})")
    if not isinstance(obj.get("refused"), bool):
        e.append("refused must be a boolean")
    if obj.get("answer_language") not in LANGS:
        e.append(f"answer_language must be one of {sorted(LANGS)}")
    return e


def build_extractor_prompt(instruction: str, aliases: str, answer: str) -> str:
    return (f"{instruction.strip()}\n\n"
            f"### Brand alias table (canonical | surface forms)\n{aliases.strip()}\n\n"
            f"### Required JSON shape\n{EXTRACTION_SCHEMA_DOC}\n\n"
            f"### Answer to analyse\n<<<ANSWER\n{answer}\nANSWER>>>\n\n"
            f"Reply with the JSON object only.")


def stratified_subsample(rows: list[dict], n: int, seed: int) -> list[dict]:
    """Balanced across model × arm (E.2), deterministic for a given seed."""
    rng = random.Random(seed)
    buckets = defaultdict(list)
    for r in rows:
        buckets[(r["model_id"], r["arm"])].append(r)
    for b in buckets.values():
        rng.shuffle(b)
    out, keys = [], sorted(buckets)
    i = 0
    while len(out) < min(n, len(rows)):
        progressed = False
        for k in keys:
            if i < len(buckets[k]):
                out.append(buckets[k][i])
                progressed = True
                if len(out) >= min(n, len(rows)):
                    break
        if not progressed:
            break
        i += 1
    return out


def cmd_extract(a):
    cfg = Config(Path(a.models))
    cache = Cache(Path(a.runs))
    logs = Path(a.runs) / "logs"
    coded_path = Path(a.runs) / "coded" / f"{a.phase}.jsonl"
    coded = read_jsonl(coded_path)
    if not coded:
        die(f"{coded_path} not found or empty — run `code --phase {a.phase}` first")
    role = a.role if getattr(a, "role", None) else ("cross_check" if a.cross else "primary")
    ex = cfg.extractor(role)
    instr_p, alias_p = Path(a.instruction), Path(a.aliases)
    missing = [str(p) for p in (instr_p, alias_p) if not p.exists()]
    if missing and not a.allow_missing_inputs:
        die("missing frozen extraction inputs " + ", ".join(missing) +
            " (prereg F2 / E.1). These are hashed at freeze; extraction must not run without them. "
            "Pass --allow-missing-inputs only for a structural dry run.")
    instruction = instr_p.read_text(encoding="utf-8") if instr_p.exists() else \
        "[PLACEHOLDER INSTRUCTION — NOT THE FROZEN TEXT; STRUCTURAL DRY RUN ONLY]"
    aliases = alias_p.read_text(encoding="utf-8") if alias_p.exists() else "[PLACEHOLDER ALIASES]"

    pool = [r for r in coded if not r["degenerate"]]
    if getattr(a, "keys", None):
        wanted = [l.strip() for l in Path(a.keys).read_text(encoding="utf-8").splitlines() if l.strip()]
        if len(set(wanted)) != len(wanted):
            die(f"--keys {a.keys}: duplicate keys in the list")
        by_key = {r["key"]: r for r in pool}
        missing = [k for k in wanted if k not in by_key]
        if missing:
            die(f"--keys {a.keys}: {len(missing)} keys not in the coded non-degenerate pool "
                f"(first: {missing[0]})")
        rows = [by_key[k] for k in wanted]
    elif a.cross or role == "cross_check":
        rows = stratified_subsample(pool, a.subsample, a.seed)
    else:
        rows = pool
    specs = []
    for r in rows:
        recp = cache.path({"phase": r["phase"], "model_id": r["model_id"], "query_id": r["query_id"],
                           "arm": r["arm"], "rep": r["rep"], "draw": r["draw"]}, r["key"])
        if not recp.exists():
            continue
        answer = (json.loads(recp.read_text(encoding="utf-8")).get("meta") or {}).get("content") or ""
        prompt = build_extractor_prompt(instruction, aliases, answer)
        specs.append({"phase": f"extract-{role}", "role": "extractor", "model_id": ex["id"],
                      "model_name": ex.get("name", ex["id"]), "query_id": r["query_id"],
                      "arm": r["arm"], "rep": r["rep"], "draw": r["draw"],
                      "subject_model_id": r["model_id"], "subject_key": r["key"],
                      "prompt": prompt, "prompt_sha256": sha256_text(prompt)})
    say(f"\nextraction ({role}, {ex.get('name', ex['id'])}) over phase '{a.phase}': "
        f"{len(specs)} answers"
        + (f" (explicit key list {a.keys}: {len(rows)} of {len(pool)})" if getattr(a, "keys", None)
           else f" (stratified subsample of {len(pool)})" if (a.cross or role == "cross_check")
           else f" (all non-degenerate)"))
    if a.dry_run or not a.go:
        say("DRY RUN — nothing was called. Add --go to execute.")
        return
    queries = load_queries(Path(a.queries))
    preflight(a, cfg, queries, specs, "extract")
    params_for = lambda s: cfg.extractor_params(role)
    execute(cfg, cache, logs, specs, params_for, a.concurrency, a.timeout, a.retries, f"extract-{role}")

    out, bad = [], 0
    dest = Path(a.runs) / "extracted" / f"{a.phase}.{role}.jsonl"
    dest.parent.mkdir(parents=True, exist_ok=True)
    with open(dest, "w", encoding="utf-8") as f:
        for s in specs:
            k = Cache.key(s, cfg.extractor_params(role))
            p = cache.path(s, k)
            if not p.exists():
                continue
            rec = json.loads(p.read_text(encoding="utf-8"))
            content = (rec.get("meta") or {}).get("content") or ""
            obj, perr = parse_extractor_json(content)
            errs = [perr] if perr else validate_extraction(obj)
            if errs:
                bad += 1
            row = {"subject_key": s["subject_key"], "extractor": ex["id"], "role": role,
                   "query_id": s["query_id"], "arm": s["arm"], "rep": s["rep"], "draw": s["draw"],
                   "subject_model_id": s["subject_model_id"],
                   "extraction": obj if not errs else None, "schema_errors": errs or None,
                   "extractor_raw": content}
            f.write(json.dumps(row, ensure_ascii=False) + "\n")
            out.append(row)
    say(f"wrote {len(out)} extraction rows -> {dest}"
        + (f"   ({bad} failed schema validation — see schema_errors)" if bad else "   (all schema-valid)"))


# ──────────────────────────────────────────────────────────── ledger and window
def cmd_ledger(a):
    cfg = Config(Path(a.models))
    cache = Cache(Path(a.runs))
    prices = load_prices(Path(a.prices))
    names = {s["id"]: s.get("name", s["id"]) for s in cfg.subjects}
    for role, e in (cfg.extractors or {}).items():
        if isinstance(e, dict) and e.get("id"):
            names[e["id"]] = e.get("name", e["id"])
    phases = a.phases or ["smoke", "pilot", "main", "extract-primary", "extract-cross_check",
                          "extract-fallback"]
    agg = defaultdict(lambda: defaultdict(lambda: {"calls": 0, "pt": 0, "ct": 0, "rt": 0, "sec": 0.0,
                                                   "billed": 0.0, "billed_n": 0, "cached": 0}))

    def tally(label, recs):
        for rec in recs:
            u = (rec.get("meta") or {}).get("usage") or {}
            d = agg[label][rec["spec"]["model_id"]]
            d["calls"] += 1
            # Some gateways report the BILLED cost per call in usage.cost (OpenRouter does).
            # That is the actual amount charged — it already accounts for prompt caching, tiered
            # rates and provider-specific pricing — so it beats anything computed from a local
            # price table. Summed here and reported alongside, never mixed into, the computed USD.
            if isinstance(u.get("cost"), (int, float)):
                d["billed"] += float(u["cost"])
                d["billed_n"] += 1
            pd_ = u.get("prompt_tokens_details") or {}
            d["cached"] += int(pd_.get("cached_tokens") or 0)
            d["pt"] += int(u.get("prompt_tokens") or 0)
            d["ct"] += int(u.get("completion_tokens") or 0)
            d["rt"] += int(((u.get("completion_tokens_details") or {}).get("reasoning_tokens")) or 0)
            d["sec"] += float(rec.get("duration_s") or 0)

    # Records superseded by a pre-freeze parameter change (prune_superseded.py moves them out of
    # runs/responses/ so they cannot re-enter an analysis) were still PAID FOR. The ledger is the
    # paper's cost table and replaces every estimate (§11, G.3), so it must not silently drop them:
    # they are tallied under their own label and included in the total, never mixed into the phase
    # whose answers are analysed. This is also what reconciles the ledger's call count with the
    # run-window log, which counts every call actually made and is deliberately never rewritten.
    labels = []
    sup_root = Path(a.runs) / "superseded"
    for ph in phases:
        tally(ph, cache.records(ph))
        labels.append(ph)
        d = sup_root / ph
        if d.exists():
            lab = f"{ph} [superseded — paid, excluded from analysis]"
            tally(lab, (json.loads(f.read_text(encoding="utf-8")) for f in sorted(d.rglob("*.json"))))
            labels.append(lab)
    phases = labels
    if not agg:
        say("cost ledger: the cache is empty — nothing has been called yet.")
        return
    grand = grand_hi = 0.0
    lo_bad = []
    for ph in [p for p in phases if p in agg]:
        say("")
        say(f"── {ph} " + "─" * (74 - len(ph)))
        say(f"{'model':<24}{'calls':>7}{'in tok':>12}{'cached':>11}{'out tok':>11}{'reason':>10}"
            f"{'USD lo':>9}{'USD hi':>9}{'BILLED':>10}{'min':>7}")
        sub = sub_hi = 0.0
        for mid, d in sorted(agg[ph].items(), key=lambda kv: -kv[1]["ct"]):
            pr = prices.get(mid) or {}
            pin, pout = float(pr.get("in_per_1m", 0)), float(pr.get("out_per_1m", 0))
            # Two bounds, because providers do not agree on whether reasoning_tokens are counted
            # INSIDE completion_tokens (OpenAI's documented convention) or reported alongside them.
            # lo assumes inside; hi assumes on top. Where rt > ct the "inside" reading is
            # arithmetically impossible, so lo is provably wrong for that model and is marked.
            usd = d["pt"] / 1e6 * pin + d["ct"] / 1e6 * pout
            usd_hi = d["pt"] / 1e6 * pin + (d["ct"] + d["rt"]) / 1e6 * pout
            impossible = d["rt"] > d["ct"]
            if impossible:
                lo_bad.append((names.get(mid, mid), d["ct"], d["rt"]))
            sub += (usd_hi if impossible else usd)
            sub_hi += usd_hi
            billed = (f"{d['billed']:,.4f}" if d["billed_n"] == d["calls"] and d["billed"]
                      else (f"~{d['billed']:,.4f}" if d["billed_n"] else "—"))
            say(f"{names.get(mid, mid):<24}{d['calls']:>7}{d['pt']:>12,}{d['cached']:>11,}"
                f"{d['ct']:>11,}{d['rt']:>10,}{('n/a' if impossible else f'{usd:,.2f}'):>9}"
                f"{usd_hi:>9,.2f}{billed:>10}{d['sec']/60:>7.1f}"
                + ("  <- reasoning NOT inside completion" if impossible else ""))
        # column widths MUST match the per-model line above: 24,7,12,11,11,10,9,9,10,7
        S = lambda k: sum(d[k] for d in agg[ph].values())
        sub_b = S("billed")
        say(f"{'subtotal':<24}{S('calls'):>7}{S('pt'):>12,}{S('cached'):>11,}"
            f"{S('ct'):>11,}{S('rt'):>10,}{sub:>9,.2f}{sub_hi:>9,.2f}"
            f"{(f'{sub_b:,.4f}' if sub_b else '—'):>10}{S('sec')/60:>7.1f}")
        grand += sub
        grand_hi += sub_hi
    say("")
    tb = sum(d["billed"] for ph in agg for d in agg[ph].values())
    tn = sum(d["billed_n"] for ph in agg for d in agg[ph].values())
    tc = sum(d["calls"] for ph in agg for d in agg[ph].values())
    if tn:
        say(f"\nTOTAL BILLED BY THE GATEWAY: ${tb:,.4f}"
            f"   ({tn} of {tc} calls report a cost; the rest are not billed through a gateway that"
            f" reports one)")
        say("This is the ACTUAL amount charged — it already includes prompt caching and any tiered"
            "\nor provider-specific rate — and it supersedes the computed columns wherever present.")
    say(f"\nTOTAL METERED USD (computed from prices.json): {grand:,.2f} (low) .. {grand_hi:,.2f} (high)" + ("" if prices else
        "   (no prices.json — token counts are exact, USD is 0 until prices are recorded)"))
    if any(d["rt"] for ph in agg for d in agg[ph].values()):
        say("The range exists because providers differ on whether reasoning tokens are counted INSIDE"
            "\ncompletion_tokens or reported alongside them: 'low' bills completion only, 'high' bills"
            "\ncompletion + reasoning. The gateway invoice for the run window settles it; record which"
            "\nreading matched (prereg §11, G.3) rather than reporting a single unlabelled number.")
    unpriced = sorted({mid for ph in agg for mid in agg[ph]
                       if not (prices.get(mid) or {}).get("out_per_1m")} ) if prices else []
    if unpriced:
        say(f"  ! NO PRICE RECORDED for {len(unpriced)} of the models called, so they contribute 0.00"
            f" to the totals above and the cost table is INCOMPLETE: {', '.join(unpriced)}."
            f"\n    Add them to prices.json before the figures are quoted anywhere.")
    for nm, ct, rt in lo_bad:
        say(f"  ! {nm}: reasoning {rt:,} EXCEEDS completion {ct:,}, so reasoning cannot be inside it —"
            f" this model is billed at the high reading in the total above.")
    if any("superseded" in k for k in agg):
        say("The total INCLUDES calls superseded by a pre-freeze parameter change: they were paid for,\nso they belong in the cost table, and they are listed separately because their answers are excluded\nfrom every analysis. Their count is why the run-window log shows more calls than the analysis set.")
    say("This ledger is the paper's cost table and replaces every estimate (prereg §11, G.3).")


def cmd_window(a):
    p = Path(a.runs) / "logs" / "run_window.json"
    if not p.exists():
        say("no run window recorded yet (no successful call).")
        return
    w = json.loads(p.read_text(encoding="utf-8"))
    say("\nRUN WINDOW (prereg C.7 — the main run must fit inside 7 consecutive days)")
    for ph, d in sorted(w.get("phases", {}).items()):
        span = ""
        try:
            t0 = datetime.strptime(d["first_call"], "%Y-%m-%dT%H:%M:%S.%fZ")
            t1 = datetime.strptime(d["last_call"], "%Y-%m-%dT%H:%M:%S.%fZ")
            days = (t1 - t0).total_seconds() / 86400
            span = f"   span {days:.2f} days" + ("   ** EXCEEDS THE 7-DAY WINDOW **" if days > 7 else "")
        except Exception:
            pass
        say(f"  {ph:<18}{d['calls']:>7} calls   {d['first_call']} → {d['last_call']}{span}")
    exp_p = Path(a.expected_models)
    expected = json.loads(exp_p.read_text(encoding="utf-8")) if exp_p.exists() else {}
    say("\nOBSERVED MODEL STRINGS (H.4 expected values / C.7 drift check)")
    drift = []
    notes = []
    for mid, d in sorted(w.get("models", {}).items()):
        strings = sorted(d.get("model_strings", {}))
        say(f"  {mid:<26}{', '.join(strings) or '—'}")
        fps = d.get("fingerprints", {})
        for fp, meta in sorted(fps.items()):
            say(f"  {'':<26}fingerprint {fp} ×{meta['count']}"
                f"  [{meta.get('first_seen', '?')} → {meta.get('last_seen', '?')}]")
        if not fps:
            say(f"  {'':<26}fingerprint —  (this endpoint returns none)")
        if len(strings) > 1:
            drift.append((mid, strings, "more than one model string within the run"))
        if expected.get(mid) and strings and expected[mid] not in strings:
            drift.append((mid, strings, f"expected {expected[mid]!r}"))
        # C.7's rule is "if the reported model OR fingerprint string for any subject changes DURING
        # THE WINDOW". Fingerprints accumulate across phases here, and a provider rolling its
        # fingerprint between the smoke test and the main run is not a within-window change, so
        # multiplicity is judged per phase by overlap of the fingerprint's own seen-interval.
        for ph, pd in sorted(w.get("phases", {}).items()):
            lo, hi = pd.get("first_call"), pd.get("last_call")
            if not (lo and hi):
                continue
            inside = sorted(fp for fp, meta in fps.items()
                            if meta.get("first_seen", "") <= hi and meta.get("last_seen", "") >= lo)
            if len(inside) > 1:
                drift.append((mid, inside, f"more than one fingerprint inside phase {ph!r}"))
        if len(fps) > 1:
            notes.append(f"{mid}: {len(fps)} fingerprints across all phases — "
                         + ", ".join(f"{fp} ×{m['count']}" for fp, m in sorted(fps.items()))
                         + " (a roll between phases is expected and reported; only a change inside"
                           " one window is a deviation)")
        # F13: this check can only be as specific as the endpoint's own labels.
        if strings and not fps:
            req = {x for x in re.split(r"[^0-9a-z.]+", mid.lower()) if x}
            got = set().union(*({x for x in re.split(r"[^0-9a-z.]+", s.lower()) if x} for s in strings))
            if got < req:
                notes.append(f"{mid}: the endpoint's string {strings[0]!r} is LESS SPECIFIC than the"
                             f" requested id (drops {', '.join(sorted(req - got))}) and this endpoint"
                             f" returns no fingerprint — a substitution inside this model family is"
                             f" invisible to this check. Stated limitation (C.7, I.7), not a pass.")
    if a.record_expected:
        rec = {mid: sorted(d.get("model_strings", {}))[0]
               for mid, d in w.get("models", {}).items() if d.get("model_strings")}
        exp_p.write_text(json.dumps(rec, indent=1), encoding="utf-8")
        say(f"\nrecorded {len(rec)} expected model strings -> {exp_p}  (paste into prereg H.4)")
    if notes:
        say("\nNOTES (not deviations — what H.4/C.7/I.7 must state about this check's reach):")
        for s in notes:
            say(f"   - {s}")
    if drift:
        say("\n** MODEL/FINGERPRINT DRIFT — log as a deviation (H.3) and re-run analyses by period (C.7):")
        for mid, strings, why in drift:
            say(f"   {mid}: {why}; saw {strings}")
    elif expected:
        say("\nno drift against the recorded expected model strings, and no fingerprint change"
            " inside any phase.")


# ───────────────────────────────────────────── release redaction (M5) and aliases (F2)
GATEWAY_REDACTED = "<gateway URL redacted for double-anonymous review — VENUE-PLAN §1.4>"


def cmd_release(a):
    """Copy runs/ to a release folder with ONLY the gateway URL redacted (request.url in every
    record; gateway_base in models.public.json). Nothing else is altered — H.5 / VENUE-PLAN §1.4."""
    import shutil
    src, dst = Path(a.runs), Path(a.out)
    if not src.exists():
        die(f"{src} does not exist")
    if dst.exists() and any(dst.iterdir()) and not a.force:
        die(f"{dst} is not empty — pass --force to overwrite")
    n_rec = 0
    for sub in ("responses", "errors"):
        for pth in (src / sub).rglob("*.json") if (src / sub).exists() else []:
            rec = json.loads(pth.read_text(encoding="utf-8"))
            if isinstance(rec.get("request"), dict) and "url" in rec["request"]:
                rec["request"]["url"] = GATEWAY_REDACTED
            out = dst / pth.relative_to(src)
            out.parent.mkdir(parents=True, exist_ok=True)
            out.write_text(json.dumps(rec, ensure_ascii=False, indent=1), encoding="utf-8")
            n_rec += 1
    for sub in ("coded", "extracted", "logs"):
        if (src / sub).exists():
            shutil.copytree(src / sub, dst / sub, dirs_exist_ok=True)
    mj = Path(a.models)
    if mj.exists():
        m = json.loads(mj.read_text(encoding="utf-8"))
        m["gateway_base"] = GATEWAY_REDACTED
        (dst / "models.public.json").write_text(json.dumps(m, indent=2, ensure_ascii=False),
                                                 encoding="utf-8")
    (dst / "RELEASE-NOTE.txt").write_text(
        f"Generated {now_iso()} by n1_pipeline.py release.\n"
        f"Redacted: request.url in {n_rec} records and gateway_base in models.public.json.\n"
        f"Everything else is byte-for-byte the working data. Raters appear only as R1-R3.\n",
        encoding="utf-8")
    say(f"release written -> {dst}   ({n_rec} records with request.url redacted)")


def cmd_aliases(a):
    """Export the brand alias table from the kit's BrandAliases-Starter sheet.
      --starter : every row, labelled starter-unverified (smoke test only)
      --frozen  : rows with verified = YES only (the F2 artifact, hashed in H.4)"""
    try:
        import openpyxl
    except ImportError:
        die("pip install openpyxl")
    kit = Path(a.kit)
    if not kit.exists():
        die(f"{kit} not found")
    wb = openpyxl.load_workbook(kit, data_only=True)
    if "BrandAliases-Starter" not in wb.sheetnames:
        die("kit has no BrandAliases-Starter sheet")
    ws = wb["BrandAliases-Starter"]
    rows, dropped = [], 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0] or "STARTER" in str(r[0]).upper():
            continue                                  # the note row
        cid, disp, cls, al, ver = (str(x).strip() if x is not None else "" for x in r[:5])
        if a.frozen and ver.upper() != "YES":
            dropped += 1
            continue
        rows.append({"canonical_id": cid, "display_name": disp, "class": cls.lower(),
                     "aliases": al, "source": "verified" if ver.upper() == "YES" else "starter-unverified"})
    if not rows:
        die("no alias rows to export" + (" — no row has verified = YES yet (F2)" if a.frozen else ""))
    bad_cls = [r["canonical_id"] for r in rows if r["class"] not in ("local", "global", "ambiguous")]
    if bad_cls:
        die(f"class must be local/global/ambiguous (Appendix C.5); offending: {bad_cls[:6]}")
    dest = Path(a.out or ("brand_aliases.csv" if a.frozen else "brand_aliases.starter.csv"))
    with open(dest, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["canonical_id", "display_name", "class", "aliases", "source"])
        w.writeheader()
        w.writerows(rows)
    digest = sha256_file(dest)
    kind = "FROZEN (verified rows only)" if a.frozen else "STARTER (unverified — smoke test only)"
    say(f"{kind}: {len(rows)} brands -> {dest}" + (f"   ({dropped} unverified rows dropped)" if a.frozen else ""))
    say(f"sha256 {digest}" + ("   (record in prereg H.4 / F2)" if a.frozen else ""))
    cls = Counter(r["class"] for r in rows)
    say("classes: " + ", ".join(f"{k}={v}" for k, v in sorted(cls.items())))


# ─────────────────────────────────────────────────────────────────── self test
def cmd_selftest(a):
    """Offline end-to-end test: synthetic cached responses -> coding -> redraw plan -> extraction
    parsing -> finalize. Makes no network call and touches no real study data."""
    import tempfile, shutil
    from argparse import Namespace
    global LIVE_ALLOWED, http_chat
    tmp = Path(tempfile.mkdtemp(prefix="n1-selftest-"))
    fails = []

    def check(cond, label):
        say(f"  [{'PASS' if cond else 'FAIL'}] {label}")
        if not cond:
            fails.append(label)

    say("\nSELFTEST (offline)\n" + "-" * 60)
    cfg = Config(Path(a.models))
    prm = {"temperature": 1.0, "max_tokens": 2000}

    # ---- guard: no live call is possible while LIVE_ALLOWED is False
    LIVE_ALLOWED = False
    try:
        http_chat(cfg, {"model": "x", "messages": []}, 1, "subject")
        check(False, "http_chat refuses to run in a dry run")
    except RuntimeError:
        check(True, "http_chat refuses to run in a dry run (LIVE_ALLOWED guard)")

    # ---- B1: validator gate
    check(validator_clean(0, "...\n0 FAIL · 0 WARN\n"), "B1: validator_clean accepts exit 0 + '0 FAIL'")
    check(not validator_clean(1, "...\n10 FAIL · 2 WARN\n"), "B1: '10 FAIL' with exit 1 is rejected")
    check(not validator_clean(0, "...\n10 FAIL · 2 WARN\n"), "B1: '10 FAIL' is rejected even with exit 0 (substring trap)")
    check(not validator_clean(1, "0 FAIL · 0 WARN"), "B1: exit 1 is rejected even if the text says 0 FAIL")

    # ---- M4: the one body builder — no system role, exact key set
    sp_subj = {"phase": "t", "role": "subject", "model_id": "m", "query_id": "Q001", "arm": "bn",
               "rep": 1, "draw": 1, "prompt": "ক", "prompt_sha256": sha256_text("ক")}
    body = build_body(cfg, sp_subj, cfg.subject_params(cfg.subjects[0]))
    check(set(body) == {"model", "messages", "temperature", "max_tokens"},
          "M4: subject body has exactly {model, messages, temperature, max_tokens}")
    check(len(body["messages"]) == 1 and body["messages"][0]["role"] == "user"
          and all(m["role"] != "system" for m in body["messages"]),
          "M4: exactly one user message and no system role (B.3)")
    body_x = build_body(cfg, {**sp_subj, "role": "extractor"}, cfg.extractor_params())
    # The invariant M4 protects is that the body is {model, messages} + EXACTLY the parameters
    # models.json declares (+ response_format for extractors) and nothing the code invents. Tying
    # the assertion to the config is stronger than a hardcoded key list: it still fails if code
    # adds a key, but it does not fail merely because a parameter was configured.
    check(set(body_x) == {"model", "messages", "response_format"} | set(cfg.extractor_params())
          and body_x["temperature"] == 0.0,
          "M4: extractor body = model + messages + response_format + exactly the models.json "
          "extractor params, at temperature 0")
    # Guard the thing that would actually corrupt the study: extractor-only settings (reasoning,
    # provider, a raised max_tokens) must NEVER leak into a subject call, because subject sampling
    # parameters are what the paper reports auditing (C.2, B.3).
    leak = set(cfg.subject_params(cfg.subjects[0])) - {"temperature", "max_tokens"}
    check(not leak and set(body) == {"model", "messages", "temperature", "max_tokens"},
          f"M4: extractor settings do not leak into subject calls (subject params stay "
          f"{{temperature, max_tokens}}; leaked: {sorted(leak) or 'none'})")
    doc_leak = [k for k in cfg.extractor_params()
                if k in ("note", "notes", "comment", "description", "rationale", "doc", "docs")
                or k.startswith("_") or isinstance(cfg.extractor_params()[k], str) and k != "model"]
    check(not doc_leak,
          f"no documentation key reaches the extractor request body or the cache key "
          f"(offending: {doc_leak or 'none'})")
    check(cfg.route("subject")[0] == cfg.base and cfg.route("extractor")[0] == cfg.extractor_base,
          f"routes resolve per role (subject -> {cfg.base}, extractor -> {cfg.extractor_base})")

    # ---- cache key behaviour
    k1 = Cache.key(sp_subj, prm)
    check(k1 == Cache.key(sp_subj, prm), "cache key is stable for identical inputs (resumable, re-runs are free)")
    check(k1 != Cache.key({**sp_subj, "prompt_sha256": sha256_text("খ")}, prm),
          "editing the query text changes the cache key (no stale reuse)")
    check(k1 != Cache.key(sp_subj, {**prm, "temperature": 0.0}), "changing a sampling parameter changes the cache key")

    # ---- B2: HTTP 200 with an error payload / no choices must NOT be cached as a success
    real_http = http_chat
    fake_cache = Cache(tmp / "b2")
    fake_logs = tmp / "b2" / "logs"
    LIVE_ALLOWED = True                       # do_one_call would otherwise refuse; http is faked
    try:
        http_chat = lambda cfg_, body_, timeout_, role_="subject": (200, {"error": {"message": "upstream timeout", "code": 502}}, "")
        r1 = do_one_call(cfg, fake_cache, fake_logs, sp_subj, prm, 5, 0)
        http_chat = lambda cfg_, body_, timeout_, role_="subject": (200, {}, "")
        r2 = do_one_call(cfg, fake_cache, fake_logs, {**sp_subj, "query_id": "Q002"}, prm, 5, 0)
        http_chat = lambda cfg_, body_, timeout_, role_="subject": (200, {"id": "x", "model": "m-2026-09", "choices": [
            {"message": {"role": "assistant", "content": "ওয়ালটন ভালো।"}, "finish_reason": "stop"}],
            "usage": {"prompt_tokens": 10, "completion_tokens": 5}}, "")
        r3 = do_one_call(cfg, fake_cache, fake_logs, {**sp_subj, "query_id": "Q003"}, prm, 5, 0)
    finally:
        http_chat = real_http
        LIVE_ALLOWED = False
    check(r1 == "error" and not fake_cache.has(sp_subj, Cache.key(sp_subj, prm))
          and fake_cache.path(sp_subj, Cache.key(sp_subj, prm), err=True).exists(),
          "B2: HTTP 200 + {error:...} is stored under errors/ and will retry, never cached as success")
    check(r2 == "error", "B2: HTTP 200 with no choices is an error, not an empty answer")
    sp3 = {**sp_subj, "query_id": "Q003"}
    check(r3 == "ok" and fake_cache.has(sp3, Cache.key(sp3, prm)),
          "B2: a genuine 200 with choices[0].message is cached as a success")
    rec3 = json.loads(fake_cache.path(sp3, Cache.key(sp3, prm)).read_text(encoding="utf-8"))
    check(rec3["meta"]["model_string"] == "m-2026-09" and rec3["meta"]["finish_reason"] == "stop"
          and rec3["meta"]["usage"]["completion_tokens"] == 5 and "url" in rec3["request"],
          "record stores model string, finish_reason, usage and the request (C.2/C.7)")
    check(not fake_cache.path(sp_subj, Cache.key(sp_subj, prm), err=True).read_text().count("sk-"),
          "no API key material in a stored record")

    # ---- deterministic coding (code stage)
    cache = Cache(tmp)
    cases = [
        ("Q001", "bn", "স্যামসাং আর ওয়ালটন দুটোই ভালো। দাম ২০,০০০ টাকার কাছাকাছি।", "stop", "valid"),
        ("Q002", "bn", "Samsung and Walton are both good options under 20,000 taka.", "stop", "language_reversion"),
        ("Q003", "bn", "   ", "stop", "degenerate"),
        ("Q004", "bn", "ভালো ভালো ভালো ভালো ভালো ভালো ভালো ভালো ভালো ভালো ভালো ভালো ", "length", "degenerate"),
        ("Q005", "bn", "আমি দুঃখিত, আমি এই বিষয়ে সাহায্য করতে পারছি না।", "stop", "refusal_candidate"),
        ("Q006", "en", "I would suggest looking at Walton or Samsung in that budget.", "stop", "pending_extractor"),
        ("Q007", "bl", "Bhai apni Walton dekhte paren, dam kom.", "stop", "pending_extractor"),
        ("Q008", "bn", "ভালো অপশন: Samsung Galaxy A15 5G (8/128GB), Walton Primo, Xiaomi Redmi Note 13.", "stop", "valid"),
        ("Q009", "bn", "Walton is a good local choice for most buyers and its service network is wide. ওয়ালটন দেশি ব্র্যান্ড হিসেবে ভালো এবং সার্ভিস সহজে পাওয়া যায়। Samsung offers better cameras. স্যামসাং ক্যামেরায় এগিয়ে।", "stop", "pending_extractor"),
    ]
    for qid, arm, content, fr, _ in cases:
        sp = {"phase": "selftest", "role": "subject", "model_id": "test-model", "query_id": qid,
              "arm": arm, "rep": 1, "draw": 1, "prompt": "x", "prompt_sha256": sha256_text("x")}
        k = Cache.key(sp, prm)
        cache.write(sp, k, {"schema": SCHEMA, "key": k, "spec": sp, "params": prm,
                            "t_request": now_iso(), "t_response": now_iso(), "duration_s": 1.0,
                            "http_status": 200, "attempts": [], "error": None,
                            "response": {"model": "test-model-2026-09"},
                            "meta": {"finish_reason": fr, "usage": {"prompt_tokens": 30, "completion_tokens": 200},
                                     "model_string": "test-model-2026-09", "fingerprint": None,
                                     "content": content, "content_chars": len(content)}})
    real = Path(a.refusal_rules)
    if real.exists():
        try:
            rdet = RefusalDetector(real, allow_missing=False)
            check(len(rdet.rules) > 0, f"F4: real {real.name} loads — {len(rdet.rules)} patterns compile (sha256 {rdet.sha256[:12]}…)")
            f_bn = rdet.check("আমি দুঃখিত, আমি এই বিষয়ে সাহায্য করতে পারছি না।")[0]
            f_en = rdet.check("I'm sorry, but I cannot recommend specific products.")[0]
            f_bl = rdet.check("Sorry bhai, ami ei bishoy e help korte parchi na.")[0]
            check(f_bn and f_en and f_bl, "F4: real rules fire on a Bangla, an English and a Banglish refusal")
            fp = [t for t in ("স্যামসাং আর ওয়ালটন দুটোই ভালো।", "I would suggest looking at Walton or Samsung.",
                              "Bhai apni Walton dekhte paren, dam kom.") if rdet.check(t)[0]]
            say(f"  [INFO] F4: real rules on three plain recommendations -> {len(fp)} fire(s)"
                + (f" {fp}" if fp else "") + "  (a fire here only costs a hedged flag under D.5)")
        except SystemExit:
            check(False, f"F4: real {real.name} has an invalid regex — see message above")
    else:
        say(f"  [INFO] F4: {real} not present — synthetic rules used below")
    rules = tmp / "refusal_rules.md"
    rules.write_text("# selftest rules (synthetic, NOT the study's F4 file)\nদুঃখিত\n^I('m| am) sorry\ncannot help\n", encoding="utf-8")
    ns = Namespace(runs=str(tmp), phase="selftest", refusal_rules=str(rules), allow_missing_refusal_rules=False)
    cmd_code(ns)
    coded = {r["query_id"]: r for r in read_jsonl(tmp / "coded" / "selftest.jsonl")}
    for qid, arm, _c, _f, expected in cases:
        got = coded.get(qid, {}).get("outcome")
        check(got == expected, f"code: {qid} [{arm}] -> '{got}' (expected '{expected}')")
    check(coded["Q008"]["script_class"] == "bn" and coded["Q008"]["bangla_ratio"] < 0.2,
          f"D.4 prose ratio: Bangla spec list with Latin product names is 'bn' (chars {coded['Q008']['bangla_ratio']:.2f}, prose {coded['Q008']['prose_ratio']:.2f})")
    check(coded["Q009"]["script_class"] == "mixed", f"D.4 prose ratio: alternating EN/BN paragraphs are 'mixed' (prose {coded['Q009']['prose_ratio']:.2f}) and go to the extractor")
    # the six realistic shapes from the D.4 decision
    shapes = [
        ("Bangla prose + Latin product names -> bn", "এই বাজেটে Samsung Galaxy A15 5G আর Walton Primo HM7 দুটোই ভালো অপশন। Samsung এর ক্যামেরা ভালো, Walton এর সার্ভিস সেন্টার বেশি।", "bn"),
        ("English -> latin", "Walton is a good local choice for most buyers in Bangladesh, and Samsung offers better cameras in the same range.", "latin"),
        ("Banglish -> latin", "Bhai ei budget e Walton dekhte paren, dam kom ar service center o beshi. Samsung er camera valo.", "latin"),
        ("Bangla with lowercase Latin loanwords -> bn", "software support আর warranty দেখলে স্যামসাং এগিয়ে, কিন্তু দাম বিবেচনায় ওয়ালটন ভালো।", "bn"),
        ("Bangla digits are dropped like Western ones -> bn", "দাম ৩৮,০০০ টাকা, Samsung Galaxy A15 5G ভালো চলে।", "bn"),
        ("parenthesised Latin is stripped -> bn", "ওয়ালটন (Walton Digi-Tech Industries Ltd) দেশি ব্র্যান্ড।", "bn"),
    ]
    for label, text, exp in shapes:
        got = script_class(text)
        check(got == exp, f"D.4 shape: {label} (prose {prose_ratio(text):.2f}, chars {bangla_ratio(text):.2f})" + ("" if got == exp else f"  [got {got}]"))
    check(coded["Q004"]["truncated"] is True, "finish_reason=length is flagged as truncated, not silently dropped (D.6)")
    rp = json.loads((tmp / "coded" / "selftest.redraw_plan.json").read_text())
    check(len(rp["redraws"]) == 2 and all(r["to_draw"] == 2 for r in rp["redraws"]),
          "redraw plan targets exactly the 2 degenerate cells at draw 2 (D.6, max 2 redraws)")
    sp = {"phase": "selftest", "role": "subject", "model_id": "test-model", "query_id": "Q003",
          "arm": "bn", "rep": 1, "draw": MAX_DRAWS, "prompt": "x", "prompt_sha256": sha256_text("x")}
    k = Cache.key(sp, prm)
    cache.write(sp, k, {"schema": SCHEMA, "key": k, "spec": sp, "params": prm, "t_request": now_iso(),
                        "t_response": now_iso(), "duration_s": 1.0, "http_status": 200, "attempts": [],
                        "error": None, "response": {}, "meta": {"finish_reason": "stop", "usage": {},
                        "model_string": "m", "fingerprint": None, "content": "", "content_chars": 0}})
    cmd_code(ns)
    rp = json.loads((tmp / "coded" / "selftest.redraw_plan.json").read_text())
    check(any(x["query_id"] == "Q003" for x in rp["persistently_degenerate"]),
          f"a cell degenerate at draw {MAX_DRAWS} is reported short, not redrawn again")

    # ---- M2: tolerant extractor-output parsing
    good = {"brands": ["Walton"], "prices": [{"amount": 20000, "currency": "BDT"}],
            "retailers": ["Daraz"], "refused": False, "answer_language": "bn"}
    gj = json.dumps(good)
    check(parse_extractor_json(gj)[0] == good, "M2: plain JSON parses")
    check(parse_extractor_json("```json\n" + gj + "\n```")[0] == good, "M2: ```json fenced output parses")
    check(parse_extractor_json("Here is the JSON you asked for:\n" + gj + "\nHope this helps.")[0] == good,
          "M2: a prose-wrapped object parses (first balanced {...} block)")
    check(parse_extractor_json("no json here")[0] is None and parse_extractor_json("{broken")[0] is None,
          "M2: garbage is rejected with an error, not crashed on")

    # ---- extraction schema validator
    check(validate_extraction(good) == [], "schema validator accepts a well-formed extraction")
    check(validate_extraction({**good, "answer_language": "bangla"}), "schema validator rejects an out-of-enum answer_language")
    check(validate_extraction({**good, "prices": [{"amount": "20k", "currency": "BDT"}]}), "schema validator rejects a non-numeric price amount")
    check(validate_extraction({**good, "refused": "no"}), "schema validator rejects a non-boolean refused")
    check(validate_extraction({**good, "recommended": ["Walton"]}) == [], "schema: recommended[] as a subset of brands is accepted")
    check(validate_extraction({**good, "recommended": ["Samsung"]}), "schema: recommended[] outside brands is rejected")
    check(validate_extraction({**good, "recommended": ["walton"]}) == [],
          "schema: recommended[] differing from brands ONLY by letter case is accepted (v0.35)")
    check(validate_extraction({**good, "recommended": ["WALTON"]}) == [],
          "schema: casefold comparison, not exact case ('WALTON' vs 'Walton')")
    check(validate_extraction({**good, "brands": ["Estée Lauder"], "recommended": ["Estee Lauder"]}),
          "schema: an ACCENT difference still fails — that is an F2 problem, not letter case")
    check(validate_extraction({**good, "recommended": "Walton"}), "schema: recommended must be a list")
    check(validate_extraction(good) == [], "schema: recommended[] is optional (absent is fine)")

    # ---- B3 + B4: finalize
    def crow(qid, arm, sclass, f4, pending, outcome, ratio=None):
        return {"key": f"k-{qid}", "phase": "ft", "model_id": "test-model", "query_id": qid, "arm": arm,
                "rep": 1, "draw": 1, "finish_reason": "stop", "truncated": False, "content_chars": 50,
                "bangla_ratio": ratio if ratio is not None else (0.95 if sclass == "bn" else 0.4 if sclass == "mixed" else 0.0),
                "script_class": sclass, "f4_fires": f4, "f4_rule": ("r" if f4 else None), "degenerate": False,
                "degenerate_reason": None, "answer_language": ("bn" if sclass == "bn" else None),
                "pending_extractor": pending, "outcome": outcome, "model_string": "m", "t_request": now_iso()}
    def erow(qid, lang, brands=("Walton",), retailers=(), refused=False, schema_fail=False, prices=()):
        ex = None if schema_fail else {"brands": list(brands), "prices": list(prices), "retailers": list(retailers),
                                        "refused": refused, "answer_language": lang}
        return {"subject_key": f"k-{qid}", "extractor": "qx", "role": "primary", "query_id": qid, "arm": "x",
                "rep": 1, "draw": 1, "subject_model_id": "test-model", "extraction": ex,
                "schema_errors": (["bad"] if schema_fail else None), "extractor_raw": "..."}
    fcases = [   # (coded row, extraction row, expected outcome, expected hedged, expected adjudicate)
        (crow("F01", "bl", "latin", False, True, "pending_extractor"), erow("F01", "banglish"), "valid", False, False),
        (crow("F02", "bl", "latin", False, True, "pending_extractor"), erow("F02", "en"), "language_reversion", False, False),
        (crow("F03", "en", "latin", False, True, "pending_extractor"), erow("F03", "en"), "valid", False, False),
        (crow("F04", "en", "latin", False, True, "pending_extractor"), erow("F04", "banglish"), "language_reversion", False, False),
        (crow("F05", "bn", "bn", True, False, "refusal_candidate"), erow("F05", "bn", brands=("Walton", "Samsung")), "valid", True, False),
        (crow("F06", "bn", "bn", True, False, "refusal_candidate"), erow("F06", "bn", brands=(), retailers=(), refused=True), "refusal", False, False),
        (crow("F07", "bn", "bn", False, False, "valid"), erow("F07", "en"), "valid", False, True),
        (crow("F08", "bn", "mixed", False, True, "pending_extractor"), erow("F08", "bn"), "valid", False, False),
        (crow("F09", "bl", "latin", False, True, "pending_extractor"), erow("F09", "bn"), "pending_adjudication", False, True),
        (crow("F10", "bl", "latin", False, True, "pending_extractor"), None, "pending_extractor", False, False),
        (crow("F11", ARM_TRANSLIT, "latin", False, True, "pending_extractor"), erow("F11", "banglish"), "valid", False, False),
        (crow("F12", "bn", "latin", False, False, "language_reversion"), erow("F12", "en"), "language_reversion", False, False),
        (crow("F13", "en", "latin", True, True, "refusal_candidate"), erow("F13", "en", brands=(), retailers=("Daraz",)), "valid", True, False),
        (crow("F14", "bl", "latin", False, True, "pending_extractor"), erow("F14", "en", schema_fail=True), "pending_extractor", False, False),
        (crow("F15", "bn", "bn", True, False, "refusal_candidate"), erow("F15", "bn", brands=(), retailers=(), prices=({"amount": 20000, "currency": "BDT"},)), "valid", True, False),
    ]
    (tmp / "coded").mkdir(exist_ok=True); (tmp / "extracted").mkdir(exist_ok=True)
    with open(tmp / "coded" / "ft.jsonl", "w", encoding="utf-8") as f:
        for c, _e, *_ in fcases:
            f.write(json.dumps(c) + "\n")
    with open(tmp / "extracted" / "ft.primary.jsonl", "w", encoding="utf-8") as f:
        for _c, e, *_ in fcases:
            if e:
                f.write(json.dumps(e) + "\n")
    cmd_finalize(Namespace(runs=str(tmp), phase="ft"))
    fin = {r["query_id"]: r for r in read_jsonl(tmp / "coded" / "ft.final.jsonl")}
    adj = {r["query_id"] for r in read_jsonl(tmp / "coded" / "ft.adjudicate.jsonl")}
    labels = {"F01": "bl+banglish -> valid", "F02": "bl+en -> reversion", "F03": "en+en -> valid",
              "F04": "en+banglish -> reversion", "F05": "B4: F4 fired but brands present -> valid + hedged",
              "F06": "B4: F4 fired and no brands/retailers -> refusal", "F07": "D.4: bn script but extractor says en -> bn wins, adjudicate",
              "F08": "M1: mixed script, extractor says bn -> valid", "F09": "D.4: latin script but extractor says bn -> pending_adjudication",
              "F10": "missing extraction -> stays pending_extractor", "F11": "bl_translit+banglish -> valid",
              "F12": "bn arm, Latin answer -> reversion regardless of extractor", "F13": "B4: hedged with retailers only -> content kept",
              "F14": "schema-failed extraction -> stays pending_extractor",
              "F15": "D.5 conjunct incl. prices: F4 fired, a price but no brand -> content kept + hedged, not refusal"}
    for c, _e, exp_out, exp_hedged, exp_adj in fcases:
        q = c["query_id"]; r = fin.get(q, {})
        ok = r.get("outcome") == exp_out and bool(r.get("hedged")) == exp_hedged and (q in adj) == exp_adj
        check(ok, f"finalize {q}: {labels[q]}" + ("" if ok else f"  [got outcome={r.get('outcome')} hedged={r.get('hedged')} adj={q in adj}]"))

    # ---- stratified subsample balance (E.2)
    rows = [{"model_id": f"m{i%3}", "arm": ARMS[i % 3], "key": str(i)} for i in range(90)]
    sub = stratified_subsample(rows, 9, 1)
    bal = Counter((r["model_id"], r["arm"]) for r in sub)
    check(len(sub) == 9 and max(bal.values()) - min(bal.values()) <= 1, "cross-extractor subsample is balanced across model x arm")
    check([r["key"] for r in stratified_subsample(rows, 9, 1)] == [r["key"] for r in sub], "subsample is deterministic for a fixed seed")

    # ---- plan shape
    qs = [{"query_id": f"Q{i:03d}", "category": "c", "en_text": "e", "bn_text": "b", "bl_text": "l", "bl_translit": ""} for i in range(1, 11)]
    plan = build_subject_plan(cfg, qs, ARMS, 2, "smoke", 7)
    check(len(plan) == 10 * 3 * len(cfg.subjects) * 2, f"smoke plan size = 10 rows x 3 arms x {len(cfg.subjects)} subjects x 2 reps = {len(plan)}")
    check(all(s_["arm"] != ARM_TRANSLIT for s_ in plan), "rows without bl_translit produce no translit calls")
    check([s_["query_id"] for s_ in plan] != sorted(s_["query_id"] for s_ in plan), "call order is randomised (B.5)")
    check([(x['query_id'], x['arm'], x['model_id'], x['rep']) for x in build_subject_plan(cfg, qs, ARMS, 2, "smoke", 7)] ==
          [(x['query_id'], x['arm'], x['model_id'], x['rep']) for x in plan], "call order is reproducible for a fixed --seed")

    # ---- models.json is honoured, nothing hardcoded
    overrides = [s_ for s_ in cfg.subjects if "temperature" in s_]
    check(all(cfg.subject_params(s_)["temperature"] == float(s_["temperature"]) for s_ in overrides) and overrides,
          f"per-model temperature override from models.json is honoured ({', '.join(s_['id'] for s_ in overrides) or 'none present'})")
    check(all(cfg.subject_params(s_)["max_tokens"] == int(cfg.subject_defaults["max_tokens"]) for s_ in cfg.subjects if "max_tokens" not in s_),
          "max_tokens comes from models.json defaults")

    shutil.rmtree(tmp, ignore_errors=True)
    say("-" * 60)
    say(f"selftest: {'ALL PASS' if not fails else str(len(fails)) + ' FAILURE(S)'}  ({len(fails)} failed)")
    if fails:
        for f in fails:
            say(f"  FAILED: {f}")
        sys.exit(1)


# ─────────────────────────────────────────────────────────────────────── CLI
def main(argv=None):
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--models", default="models.json", help="single source of truth (default: models.json)")
    common.add_argument("--kit", default="QUERY-AUTHORING-KIT.xlsx")
    common.add_argument("--queries", default="queries.csv")
    common.add_argument("--runs", default="runs", help="cache + logs root (default: runs/)")
    common.add_argument("--prices", default="prices.json",
                        help='optional {"<model id>": {"in_per_1m": x, "out_per_1m": y}}')
    common.add_argument("--validator", default="validate_queries.py")
    common.add_argument("--refusal-rules", default="refusal_rules.md", help="prereg F4 (one regex per line)")
    common.add_argument("--instruction", default="extractor_instruction.md",
                        help="frozen extraction instruction (E.1)")
    common.add_argument("--aliases", default="brand_aliases.csv", help="frozen alias table (F2)")
    common.add_argument("--expected-models", default="expected_models.json",
                        help="H.4 expected model strings")
    common.add_argument("--seed", type=int, default=20260902, help="call-order / subsample seed (B.5)")
    common.add_argument("--concurrency", type=int, default=4)
    common.add_argument("--timeout", type=int, default=180)
    common.add_argument("--retries", type=int, default=4,
                        help="retries on 429/5xx/timeout, exponential backoff with jitter")
    common.add_argument("--max-calls", type=int, default=40000, help="refuse a plan larger than this")
    common.add_argument("--dry-run", action="store_true", help="print the plan and stop (the default)")
    common.add_argument("--go", action="store_true",
                        help="THE SPEND GATE: allow live calls, only after preflight passes")

    ap = argparse.ArgumentParser(
        prog="n1_pipeline.py",
        description="N1 study pipeline. Dry by default: no network call happens without --go.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="Typical order:  export -> run --smoke --dry-run -> (authors say go) -> "
               "run --smoke --go -> code -> extract -> ledger -> window --record-expected\n"
               "All flags go AFTER the subcommand:  n1_pipeline.py run --smoke --go")
    sub = ap.add_subparsers(dest="cmd", required=True)

    sub.add_parser("export", parents=[common], help="kit -> queries.csv (FINAL rows only) + sha256")

    for name, helptext in (("plan", "print the call plan + cost estimate (never calls)"),
                           ("run", "execute subject calls (needs --go)")):
        p = sub.add_parser(name, parents=[common], help=helptext)
        g = p.add_mutually_exclusive_group()
        g.add_argument("--smoke", action="store_true",
                       help="10 smoke_test rows x 3 arms (+ bl_translit where present) x subjects x 2 reps (H.2)")
        g.add_argument("--pilot", action="store_true",
                       help="6 calibration rows x 3 arms x subjects x 3 reps (G.1)")
        g.add_argument("--main", action="store_true",
                       help="all FINAL queries x 3 arms x subjects x r reps (C.4)")
        g.add_argument("--translit", action="store_true",
                       help="robust50 rows, bl_translit arm only (A.5b, F.10)")
        p.add_argument("--reps", type=int, default=0, help="override the phase default")
        p.add_argument("--no-translit", action="store_true", help="smoke: skip the bl_translit arm")

    p = sub.add_parser("code", parents=[common], help="deterministic outcome coding + redraw plan (offline)")
    p.add_argument("--phase", default="smoke")
    p.add_argument("--allow-missing-refusal-rules", action="store_true",
                   help="structural coding only; 'refused' stays null and outcomes are incomplete")

    p = sub.add_parser("redraw", parents=[common], help="execute the redraw plan from `code` (needs --go)")
    p.add_argument("--phase", default="smoke")

    p = sub.add_parser("finalize", parents=[common],
                       help="fold the primary extraction into final outcome codes (offline; B3/B4)")
    p.add_argument("--phase", default="smoke")

    p = sub.add_parser("release", parents=[common],
                       help="copy runs/ with the gateway URL redacted (H.5, VENUE-PLAN 1.4)")
    p.add_argument("--out", default="release")
    p.add_argument("--force", action="store_true")

    p = sub.add_parser("aliases", parents=[common],
                       help="export brand_aliases from the kit: --starter (smoke) or --frozen (verified=YES)")
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--starter", action="store_true")
    g.add_argument("--frozen", action="store_true")
    p.add_argument("--out", default=None)

    p = sub.add_parser("extract", parents=[common], help="run an extractor over coded answers (needs --go)")
    p.add_argument("--phase", default="smoke")
    g = p.add_mutually_exclusive_group()
    g.add_argument("--primary", action="store_true", help="primary extractor, all answers (E.1)")
    g.add_argument("--cross", action="store_true", help="cross-extractor, stratified subsample (E.2)")
    g.add_argument("--role", choices=["primary", "cross_check", "fallback"], default=None,
                   help="run this extractors.<role> from models.json (G.2 contest)")
    p.add_argument("--subsample", type=int, default=2000, help="cross-extractor subsample size")
    p.add_argument("--keys", default=None,
                   help="file with one subject key per line: extract exactly these answers "
                        "(G.2 contest set); overrides the subsample")
    p.add_argument("--allow-missing-inputs", action="store_true",
                   help="structural dry run without the frozen instruction / alias table")

    p = sub.add_parser("ledger", parents=[common], help="per-model token and cost ledger (offline)")
    p.add_argument("--phases", nargs="*", default=None)

    p = sub.add_parser("window", parents=[common], help="run-window + model-string drift report (offline)")
    p.add_argument("--record-expected", action="store_true",
                   help="write the observed model strings as the H.4 expected values")

    sub.add_parser("selftest", parents=[common], help="offline end-to-end test on synthetic data")

    a = ap.parse_args(argv)
    if a.cmd == "plan":
        a.dry_run, a.go = True, False
        return cmd_run(a)
    return {"export": cmd_export, "run": cmd_run, "code": cmd_code, "redraw": cmd_redraw,
            "extract": cmd_extract, "finalize": cmd_finalize, "release": cmd_release,
            "aliases": cmd_aliases, "ledger": cmd_ledger, "window": cmd_window,
            "selftest": cmd_selftest}[a.cmd](a)


if __name__ == "__main__":
    main()
