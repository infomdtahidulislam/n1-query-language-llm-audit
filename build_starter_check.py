#!/usr/bin/env python3
"""build_starter_check.py — N1: post-hoc author confirmation of the 49 starter-table classes (28 Sep 2026).

The frozen brand_aliases.csv carries the 49 starter entities with source "starter-unverified": their local/global/
ambiguous classes were AI-drafted on 1 Sep 2026 from the first author's market knowledge and, apart from the six
retailer-class confirmations of 13 Sep, were never confirmed row by row before the freeze. This builds one sheet per
author with the 49 rows in alphabetical order (no frequencies, no results — classification is by the C.1
headquarters/ownership rule, not usage), a dropdown for the class, and a note column. Each author fills it
independently; the returned files are banked as received and compared. The frozen table is never edited.

Usage: python build_starter_check.py  ->  STARTER-CHECK-Tahidul.xlsx, STARTER-CHECK-Maksuda.xlsx
"""
import csv, hashlib
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

W = Path(".")
sha = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
with open(W / "brand_aliases.starter.csv", encoding="utf-8", newline="") as fh:
    rows = list(csv.DictReader(fh))
assert len(rows) == 49 and all(r["source"] == "starter-unverified" for r in rows)
with open(W / "brand_aliases.csv", encoding="utf-8", newline="") as fh:
    frozen = {r["canonical_id"]: r for r in csv.DictReader(fh)}
for r in rows:                                   # the frozen table carries the same class for every starter entity
    assert frozen[r["canonical_id"]]["class"] == r["class"], r["canonical_id"]
CONFIRMED_13SEP = {"chaldal", "pickaboo", "rokomari", "ryans", "startech", "daraz"}   # retailer-class confirmations, v0.40/v0.4x
rows.sort(key=lambda r: r["canonical_id"])

RULE = ("C.1 rule (registered): classify by headquarters / ownership, not by where the product is made or sold. "
        "local = Bangladeshi-owned brand; global = foreign-owned (a multinational with a Bangladesh plant, assembler or "
        "joint venture is still global); ambiguous = foreign parent with a strong Bangladesh identity (subsidiary case). "
        "Judge each row on ownership only; do not use anything the models said. Fill 'your class' for every row, add a "
        "note where you disagree with the drafted class or are unsure, save, and send the file back unchanged otherwise.")

for author in ("Tahidul", "Maksuda"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "starter-check"
    ws["A1"] = f"N1 — starter-table class confirmation — {author} — 28 Sep 2026 (independent; no discussion before returning)"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = RULE
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:G2")
    ws.row_dimensions[2].height = 75
    head = ["canonical_id", "display_name", "drafted class (1 Sep 2026)", "aliases in the frozen table", "confirmed 13 Sep as retailer class?", "your class (local / global / ambiguous)", "note (why, or source of ownership)"]
    for j, h in enumerate(head, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E8E4DA"); c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, r in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=r["canonical_id"])
        ws.cell(row=i, column=2, value=r["display_name"])
        ws.cell(row=i, column=3, value=r["class"])
        ws.cell(row=i, column=4, value=frozen[r["canonical_id"]]["aliases"])
        ws.cell(row=i, column=5, value="yes" if r["canonical_id"] in CONFIRMED_13SEP else "")
        ws.cell(row=i, column=6, value="")
        ws.cell(row=i, column=7, value="")
        ws.cell(row=i, column=6).fill = PatternFill("solid", fgColor="FFF6CC")
    dv = DataValidation(type="list", formula1='"local,global,ambiguous"', allow_blank=True, showErrorMessage=True,
                        errorTitle="class", error="Choose local, global or ambiguous")
    ws.add_data_validation(dv)
    dv.add(f"F5:F{4 + len(rows)}")
    for col, w in zip("ABCDEFG", (16, 16, 22, 48, 18, 30, 44)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    out = W / f"STARTER-CHECK-{author}.xlsx"
    wb.save(out)
    print(f"wrote {out.name} sha256 {sha(out)[:16]}… ({len(rows)} rows)")
print("starter csv sha256", sha(W / "brand_aliases.starter.csv")[:16], "| frozen brand_aliases.csv sha256", sha(W / "brand_aliases.csv")[:16])
